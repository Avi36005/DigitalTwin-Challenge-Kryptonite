"""GlucoTwin architecture diagram: one spec, two renderers.

`draw_pptx(slide, spec)` adds native, editable PowerPoint shapes; `draw_png(path, spec)` renders the same layout with
matplotlib for the README and the PDFs. Coordinates are inches on a 13.333 x 7.5 in (16:9) slide, origin top-left.
`check_fit(spec)` measures every line and fails the build if text would spill out of its box.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

INK, MUTED, RULE, BLUE = "#0B1220", "#3F4652", "#C9CED6", "#2563EB"
FONT = "Arial"
TITLE_PT, LINE_PT, ZONE_PT, EDGE_PT, LEGEND_PT = 11.5, 9.5, 8.5, 8.5, 8.5
PAD, TITLE_H, LINE_H = 0.12, 0.32, 0.225


@dataclass
class Node:
    key: str
    x: float
    y: float
    w: float
    h: float
    title: str
    lines: list[str]
    accent: bool = False  # serving model: blue border
    dashed: bool = False  # research-only candidate: dashed border, never served

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def r(self) -> float:
        return self.x + self.w

    @property
    def b(self) -> float:
        return self.y + self.h


@dataclass
class Zone:
    x: float
    y: float
    w: float
    h: float
    label: str


@dataclass
class Edge:
    points: list[tuple[float, float]]
    label: str = ""
    label_at: tuple[float, float] | None = None  # left edge, vertical centre
    arrow: bool = True
    dashed: bool = False


@dataclass
class Spec:
    zones: list[Zone] = field(default_factory=list)
    nodes: dict[str, Node] = field(default_factory=dict)
    edges: list[Edge] = field(default_factory=list)
    legend: str = ""
    legend_y: float = 7.0


def node_h(n_lines: int) -> float:
    return PAD + TITLE_H + n_lines * LINE_H + PAD * 0.7


def build_spec(n_profile: int, n_dynamic: int, top: float = 1.30) -> Spec:
    s = Spec()

    def add(key, x, y, w, title, lines, accent=False, dashed=False):
        s.nodes[key] = Node(key, x, y, w, node_h(len(lines)), title, lines, accent, dashed)
        return s.nodes[key]

    gap, zl = 0.17, 0.36  # gap between stacked boxes; space for a zone label
    # data sources: the two streams the brief requires, then the synthetic and external cohorts
    sx, sw = 0.52, 2.30
    st = add("static", sx, top + zl, sw, "Static stream",
             ["baseline profile (EHR-like)", "age, sex, BMI, HbA1c", "fasting glucose, insulin,", "triglycerides, HDL"])
    dy = add("dynamic", sx, st.b + gap, sw, "Dynamic stream",
             ["Dexcom G6 CGM, native 5-min", "Fitbit heart rate, METs,", "activity minutes"])
    sy = add("synthetic", sx, dy.b + gap, sw, "Synthetic patients",
             ["5 fictional profiles (CC0)", "gaps, late meals, short history"])
    ex = add("external", sx, sy.b + gap, sw, "External cohort",
             ["ShanghaiT2DM, 98 people", "15-min CGM, research only"])

    # ingestion
    ix, iw = 3.22, 1.66
    ad = add("adapter", ix, top + zl + 0.55, iw, "Dataset adapter",
             ["SHA-256 verified", "native readings", "observed_at stamp", "available_at stamp"])
    cs = add("store", ix, ad.b + 0.55, iw, "Canonical store",
             ["Parquet per patient", "SQLite index", "quality flag per", "reading"])

    # twin: three columns shared by both lanes
    c1, c2, c3 = (5.42, 1.98), (7.70, 2.86), (10.88, 2.03)
    lane_top = top + zl
    rp = add("replay", c1[0], lane_top, c1[1], "Replay session",
             ["backend-owned clock", "releases a reading", "once available_at ≤ t", "seek / reset → new", "revision"])
    tw = add("twin", c2[0], lane_top, c2[1], "Twin engine (serving model)",
             ["1  causal features, ≤ 180 min history",
              f"2  fusion LightGBM: {n_profile} profile + {n_dynamic} dynamic",
              "3  bounded patient update (±25 mg/dL)",
              "4  80 % range · warning > 180 mg/dL",
              "5  abstain / degraded / stale guards"], accent=True)
    db = add("dashboard", c3[0], lane_top, c3[1], "Clinician dashboard",
             ["FastAPI /api/v1 → React", "warnings with outcomes", "clinician review + note", "review ledger (SQLite)",
              "evidence · data quality"])
    row1 = max(n.h for n in (rp, tw, db))
    bottom = max(ex.b, cs.b) + 0.16
    off_top = max(lane_top + row1 + 0.62, bottom - 0.16 - node_h(4) - zl)  # offline lane ends with the sources
    fb = add("features", c1[0], off_top + zl, c1[1], "Origin table",
             ["causal features", "targets +30 / +60 min", "split 60/10/10/20", "purged boundaries"])
    fx = add("experiment", c2[0], off_top + zl, c2[1], "Frozen experiment",
             ["baselines: persistence, trend, Ridge", "LightGBM · CatBoost search", "validation-only selection",
              "one final-test run"])
    ev = add("evidence", c3[0], off_top + zl, c3[1], "Evidence + registry",
             ["metrics.json, report", "model card, disclosures", "checksummed artifacts", "rollback to previous"])
    for row in ((rp, tw, db), (fb, fx, ev)):  # equal heights so each lane reads as a row
        h = max(n.h for n in row)
        for n in row:
            n.h = h

    bottom = max(bottom, ev.b + 0.16)
    s.zones = [Zone(0.40, top, 2.54, bottom - top, "DATA SOURCES"),
               Zone(3.10, top, 1.90, bottom - top, "INGESTION"),
               Zone(5.26, top, 7.83, rp.b + 0.16 - top, "ONLINE · DIGITAL TWIN (LOCAL REPLAY)"),
               Zone(5.26, off_top, 7.83, bottom - off_top, "OFFLINE · MODEL BUILD")]
    s.legend_y = bottom + 0.14

    mid_top, mid_bot = rp.y + 0.62, fb.y + 0.62
    E = s.edges
    bus1 = (sx + sw + ix) / 2  # sources → adapter through a bus
    for n in (st, dy, sy, ex):
        E.append(Edge([(n.r, n.cy), (bus1, n.cy)], arrow=False))
    E.append(Edge([(bus1, st.cy), (bus1, ex.cy)], arrow=False))
    E.append(Edge([(bus1, ad.cy), (ad.x, ad.cy)]))
    E.append(Edge([(ad.cx, ad.b), (ad.cx, cs.y)]))
    bus2 = (ix + iw + c1[0]) / 2  # store → replay (online) and → origin table (offline)
    E.append(Edge([(cs.r, cs.cy), (bus2, cs.cy)], arrow=False))
    E.append(Edge([(bus2, mid_top), (bus2, mid_bot)], arrow=False))
    E.append(Edge([(bus2, mid_top), (rp.x, mid_top)]))
    E.append(Edge([(bus2, mid_bot), (fb.x, mid_bot)]))
    E.append(Edge([(rp.r, mid_top), (tw.x, mid_top)]))
    E.append(Edge([(tw.r, mid_top), (db.x, mid_top)]))
    E.append(Edge([(fb.r, mid_bot), (fx.x, mid_bot)]))
    E.append(Edge([(fx.r, mid_bot), (ev.x, mid_bot)]))
    gap_mid = (rp.b + 0.16 + off_top) / 2  # centre of the white band between the lanes
    cand_lines = ["research only, offline;", "served only if every", "release gate passes"]
    ch = node_h(len(cand_lines))
    add("candidate", c1[0], gap_mid - ch / 2, 1.58, "Candidates", cand_lines, dashed=True)
    E.append(Edge([(fx.cx, fx.y), (fx.cx, tw.b)], "registered model, SHA-256", (fx.cx + 0.09, gap_mid)))
    E.append(Edge([(ev.cx, ev.y), (ev.cx, db.b)], "evidence page", (ev.cx + 0.09, gap_mid)))
    dx = fb.r - 0.32
    E.append(Edge([(dx, fb.y), (dx, rp.b)], "same feature + adaptation code", (dx + 0.09, gap_mid),
                  arrow=False, dashed=True))
    s.legend = ("Solid arrow: data flow   ·   Dashed line: identical code offline and online   ·   Blue: serving model   ·   "
                "Dashed box: research only, not served   ·   CPU only, local, no cloud, no LLM   ·   not a medical device")
    return s


def check_fit(spec: Spec) -> None:
    """Raise if any title or line is wider than its box (measured with the Arial metrics matplotlib uses)."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(13.333, 7.5), dpi=100)
    r = fig.canvas.get_renderer()
    bad = []
    for n in spec.nodes.values():
        for txt, pt, weight in [(n.title, TITLE_PT, "bold"), *[(ln, LINE_PT, "normal") for ln in n.lines]]:
            t = fig.text(0, 0, txt, fontsize=pt, fontweight=weight, family=FONT)
            w_in = t.get_window_extent(r).width / fig.dpi
            if w_in > n.w - 2 * PAD + 0.02:
                bad.append(f"{n.key}: '{txt}' {w_in:.2f} in > {n.w - 2 * PAD:.2f} in")
    plt.close(fig)
    if bad:
        raise ValueError("architecture text does not fit:\n" + "\n".join(bad))


# ------------------------------------------------------------------------------------------------ PowerPoint


def draw_pptx(slide, spec: Spec) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    def plain(shape):  # drop the theme style reference: no shadow, no theme line or fill
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)
        return shape

    def text(x, y, w, h, lines, size, bold=False, color=INK, italic=False, spacing=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            if spacing:
                p.line_spacing = Pt(spacing)
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, FONT
            r.font.color.rgb = rgb(color)
        return tb

    for z in spec.zones:
        sh = plain(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(z.x), Inches(z.y), Inches(z.w), Inches(z.h)))
        sh.fill.background()
        sh.line.color.rgb = rgb(RULE)
        sh.line.width = Pt(0.75)
        text(z.x + 0.12, z.y + 0.10, z.w - 0.2, 0.2, [z.label], ZONE_PT, True, MUTED)

    for e in spec.edges:  # connectors before boxes so line ends tuck under borders
        segs = list(zip(e.points[:-1], e.points[1:], strict=True))
        for i, ((x1, y1), (x2, y2)) in enumerate(segs):
            c = plain(slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)))
            c.line.color.rgb = rgb(INK)
            c.line.width = Pt(1.0)
            if e.dashed:
                c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if e.arrow and i == len(segs) - 1:
                ln = c.line._get_or_add_ln()
                ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        if e.label:
            lx, ly = e.label_at
            text(lx, ly - 0.08, 3.0, 0.18, [e.label], EDGE_PT, False, MUTED, italic=True)

    for n in spec.nodes.values():
        sh = plain(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(n.x), Inches(n.y), Inches(n.w),
                                          Inches(n.h)))
        sh.adjustments[0] = 0.05
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb("#FFFFFF")
        sh.line.color.rgb = rgb(BLUE if n.accent else MUTED if n.dashed else INK)
        sh.line.width = Pt(1.75 if n.accent else 1.0)
        if n.dashed:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        text(n.x + PAD, n.y + PAD, n.w - 2 * PAD, TITLE_H, [n.title], TITLE_PT, True)
        text(n.x + PAD, n.y + PAD + TITLE_H, n.w - 2 * PAD, len(n.lines) * LINE_H, n.lines, LINE_PT,
             spacing=LINE_H * 72)

    text(0.40, spec.legend_y, 12.6, 0.25, [spec.legend], LEGEND_PT, False, MUTED)


# ------------------------------------------------------------------------------------------------ PNG


def draw_png(path: Path, spec: Spec, title: str | None = None, subtitle: str | None = None, dpi: int = 220) -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, Rectangle

    fig = plt.figure(figsize=(13.333, 7.5))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 13.333)
    ax.set_ylim(7.5, 0)
    ax.axis("off")
    if title:
        ax.text(0.40, 0.40, title, fontsize=26, fontweight="bold", color=INK, va="top", family=FONT)
    if subtitle:
        ax.text(0.40, 0.90, subtitle, fontsize=12, color=INK, va="top", family=FONT)

    for z in spec.zones:
        ax.add_patch(Rectangle((z.x, z.y), z.w, z.h, fill=False, ec=RULE, lw=0.75))
        ax.text(z.x + 0.12, z.y + 0.10, z.label, fontsize=ZONE_PT, fontweight="bold", color=MUTED, va="top", family=FONT)

    for e in spec.edges:
        segs = list(zip(e.points[:-1], e.points[1:], strict=True))
        for i, ((x1, y1), (x2, y2)) in enumerate(segs):
            style = "-|>" if e.arrow and i == len(segs) - 1 else "-"
            ax.annotate("", xy=(x2, y2), xytext=(x1, y1), zorder=1,
                        arrowprops=dict(arrowstyle=style, color=INK, lw=1.0, shrinkA=0, shrinkB=0,
                                        linestyle=(0, (4, 3)) if e.dashed else "-", mutation_scale=11))
        if e.label:
            lx, ly = e.label_at
            ax.text(lx, ly, e.label, fontsize=EDGE_PT, color=MUTED, style="italic", va="center", family=FONT, zorder=3)

    for n in spec.nodes.values():
        r = 0.05 * min(n.w, n.h)
        ax.add_patch(FancyBboxPatch((n.x + r, n.y + r), n.w - 2 * r, n.h - 2 * r, boxstyle=f"round,pad={r}",
                                    fc="white", ec=BLUE if n.accent else MUTED if n.dashed else INK,
                                    lw=1.75 if n.accent else 1.0, ls=(0, (4, 3)) if n.dashed else "-", zorder=2))
        ax.text(n.x + PAD, n.y + PAD, n.title, fontsize=TITLE_PT, fontweight="bold", color=INK, va="top",
                family=FONT, zorder=3)
        for i, ln in enumerate(n.lines):
            ax.text(n.x + PAD, n.y + PAD + TITLE_H + i * LINE_H, ln, fontsize=LINE_PT, color=INK, va="top",
                    family=FONT, zorder=3)

    ax.text(0.40, spec.legend_y, spec.legend, fontsize=LEGEND_PT, color=MUTED, va="top", family=FONT)
    fig.savefig(path, dpi=dpi, facecolor="white")
    plt.close(fig)
    return path
