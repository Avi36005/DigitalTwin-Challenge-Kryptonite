"""Research evidence workbook: every public claim traced to a saved source value, with visible calculations.

Reads the saved artifacts and reports at build time (nothing is fitted, no assessment is re-opened), stores each source
value once in the ``Sources`` sheet at full precision with its locator and file SHA-256, and builds the claim index,
cohorts, forecast and warning tables, the study ledger, the gate calculations, the verification record and the
guardrail map from Excel formulas that reference those cells. A source file missing on this installation becomes the
text "unavailable" (never 0) and every claim that depends on it becomes ``pending``.

Writes (aggregate only; no participant rows, no private paths):
  docs/submission/GlucoTwin_Research_Workbook.xlsx
  docs/submission/GlucoTwin_Research_Workbook_Methodology.md
  docs/submission/GlucoTwin_Research_Claims.json

Adding deck or narration claims: append to ``DECK_CLAIMS`` (new claims) or ``EXTRA_LOCATIONS`` (a further place where an
existing claim appears). See the comment above ``DECK_CLAIMS``.

Usage:  uv run --group docs python scripts/build_research_workbook.py
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import math
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs/submission"
OUT_XLSX = OUT_DIR / "GlucoTwin_Research_Workbook.xlsx"
OUT_MD = OUT_DIR / "GlucoTwin_Research_Workbook_Methodology.md"
OUT_JSON = OUT_DIR / "GlucoTwin_Research_Claims.json"

FUSION = "artifacts/experiments/gt-primary-2026-10-13-fusion"
RES = "artifacts/research"
VER = "reports/verification/final_2026-10-22_145326Z"

BLUE, INK, GREY, LINE = "2563EB", "111827", "6B7280", "D1D5DB"
STATUSES = ["supported", "supported-arithmetic", "tested", "implemented", "negative-result", "author-reported",
            "proposed/illustrative", "pending", "unsupported/contradicted"]
STATUS_MEANING = {
    "supported": "The stated value or fact matches a saved artifact, report or dataset record named in Sources.",
    "supported-arithmetic": "Follows by arithmetic from stated definitions (formula shown); no experiment involved.",
    "tested": "Behaviour covered by a test that passed in the saved test record (JUnit) named in the row.",
    "implemented": "Present in code or configuration only; no saved passing test is cited for the behaviour.",
    "negative-result": "A saved study result that did not meet its declared bar; reported as negative, not as a gain.",
    "author-reported": "Stated by the team without a saved log, artifact or primary source in the repository.",
    "proposed/illustrative": "A proposal, scenario or illustration; not an evidential claim.",
    "pending": "Evidence is not on this installation or a primary-source check is still to be done.",
    "unsupported/contradicted": "The saved source disagrees with the stated value, or no support exists.",
}


# --------------------------------------------------------------------------- small helpers

def sha256(p: Path) -> str | None:
    if not p.is_file():
        return None
    h = hashlib.sha256()
    with p.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def git(*args: str) -> str:
    try:
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=False).stdout.strip()
    except OSError:
        return ""


def tracked(rel: str) -> bool:
    return bool(git("ls-files", "--", rel))


def public_status(rel: str) -> str:
    if not (ROOT / rel).exists():
        return "missing on this installation"
    if tracked(rel):
        return "public in repo (tracked)"
    return "local / reproducible-only (git-ignored)"


def first_commit(rel: str) -> tuple[str, str]:
    """(short hash, ISO date) of the commit that added ``rel``."""
    out = git("log", "--diff-filter=A", "--format=%h %aI", "--", rel).splitlines()
    return tuple(out[-1].split(" ", 1)) if out else ("", "")  # type: ignore[return-value]


def last_commit(rel: str) -> tuple[str, str]:
    out = git("log", "-1", "--format=%h %aI", "--", rel)
    return tuple(out.split(" ", 1)) if out else ("", "")  # type: ignore[return-value]


def scrub(s: str) -> str:
    """Repository-relative text only: no home directories or session paths in the public workbook."""
    s = re.sub(r"(?:/Users/[^/\s]+|~)/projects/Digitaltwin/?", "", s)
    s = re.sub(r"/Users/[^\s]+", "<local path>", s)
    return s.replace("<session-scratch>", "<scratch>")


def utc(s: str | None) -> dt.datetime | None:
    """ISO timestamp → naive UTC datetime for an Excel cell."""
    if not s:
        return None
    d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    if d.tzinfo is None:
        return d
    return d.astimezone(dt.UTC).replace(tzinfo=None)


_cache: dict[str, Any] = {}


def load(rel: str) -> Any:
    if rel not in _cache:
        p = ROOT / rel
        _cache[rel] = json.loads(p.read_text()) if p.is_file() else None
    return _cache[rel]


def resolve(doc: Any, pointer: str) -> Any:
    """JSON pointer with list selectors: ``/events[model=x&split=test&horizon=60]/precision``, ``/ci95[1]``."""
    cur = doc
    for seg in [s for s in pointer.split("/") if s]:
        if cur is None:
            return None
        m = re.fullmatch(r"(.*?)\[(.+)\]", seg)
        key, sel = (m.group(1), m.group(2)) if m else (seg, None)
        if key:
            cur = cur.get(key) if isinstance(cur, dict) else None
        if sel is not None and cur is not None:
            if sel.isdigit():
                cur = cur[int(sel)] if isinstance(cur, list) and int(sel) < len(cur) else None
            else:
                conds = [c.split("=", 1) for c in sel.split("&")]
                cur = next((x for x in cur if all(str(x.get(k)) == v for k, v in conds)), None)
    return cur


# --------------------------------------------------------------------------- source files

# id → (path, what it is, JSON pointer of its own timestamp or None, protocol / revision identity)
FILES: dict[str, tuple[str, str, str | None, str]] = {
    "registry": ("artifacts/registry.json", "Model registry (serving mapping)", None, "registry"),
    "fusion": (f"{FUSION}/metrics.json", "Serving experiment metrics (fusion LightGBM)", "/created_at",
               "configs/experiments/gt-primary-2026-10-13-fusion.yaml"),
    "fusion_pred": (f"{FUSION}/predictions.parquet", "Serving experiment saved predictions (read by scripts/warning_utility.py)",
                    None, "configs/experiments/gt-primary-2026-10-13-fusion.yaml"),
    "trend": (f"{FUSION}/error_by_trend.json", "Serving model error by glucose trend (validation segment)", "/generated_at",
              "scripts/error_by_trend.py"),
    "parity": ("reports/serving_parity.json", "Saved app-vs-evaluation parity run", "/checked_at",
               "scripts/serving_parity.py"),
    "parity_final": (f"{VER}/parity/serving_parity.json", "Parity re-run in the final verification ledger", "/checked_at",
                     "scripts/serving_parity.py"),
    "nested": (f"{RES}/gt-nested-2026-10-12-v1/report.json", "Nested participant-held-out + external study (D-032)",
               "/frozen_at", "configs/experiments/gt-nested-2026-10-12-v1.yaml"),
    "twoh": (f"{RES}/gt-2h-2026-10-20/two_hour_summary.json", "Two-hour study (D-043)", "/frozen_at",
             "configs/experiments/gt-2h-2026-10-20.yaml"),
    "twoh_ext": (f"{RES}/gt-2h-ext-2026-10-20/two_hour_external_summary.json", "External two-hour study (D-045)",
                 "/opened/opened_at", "configs/experiments/gt-2h-ext-2026-10-20.yaml"),
    "clf": (f"{RES}/gt-2h-clf-2026-10-20/summary.json", "Two-hour onset classifier (D-044)", "/opened/opened_at",
            "configs/experiments/gt-2h-clf-2026-10-20.yaml"),
    "warnrange": (f"{RES}/gt-warn-range-2026-10-20/summary.json", "Warning/range/profile re-analysis (D-046)",
                  "/opened/opened_at", "configs/experiments/gt-warn-range-2026-10-20.yaml"),
    "blend": (f"{RES}/gt-blend-2026-10-21/development.json", "Blend study development decision (D-048)", None,
              "configs/experiments/gt-blend-2026-10-21.yaml"),
    "blend_proto": (f"{RES}/gt-blend-2026-10-21/protocol_written.json", "Blend protocol-written marker (D-048)",
                    "/written_at", "configs/experiments/gt-blend-2026-10-21.yaml"),
    "retrain": (f"{RES}/gt-retrain-2026-10-21/development.json", "Retraining study development results (D-049)",
                "/started_at", "configs/experiments/gt-retrain-2026-10-21.yaml"),
    "ranges": (f"{RES}/gt-ranges-2026-10-21/summary.json", "Serving range study (D-050)", "/generated_at",
               "configs/experiments/gt-ranges-2026-10-21.yaml"),
    "tune": (f"{RES}/gt-tune-2026-10-21/summary.json", "Nested tuning study (D-051)", "/started_at",
             "configs/experiments/gt-tune-2026-10-21.yaml"),
    "joint": (f"{RES}/gt-joint-2026-10-21/summary.json", "Joint selection study (D-052)", "/started_at",
              "configs/experiments/gt-joint-2026-10-21.yaml"),
    "constrained": (f"{RES}/gt-constrained-2026-10-21/summary.json", "Constrained selection study (D-053)",
                    "/manifest/started_at", "configs/experiments/gt-constrained-2026-10-21.yaml"),
    "foundation": (f"{RES}/gt-foundation-2026-10-21/summary.json", "Zero-shot foundation baseline (D-054)",
                   "/generated_at", "configs/experiments/gt-foundation-2026-10-21.yaml"),
    "foundation_meta": (f"{RES}/gt-foundation-2026-10-21/run_meta.json", "Foundation baseline run metadata (D-054)",
                        None, "configs/experiments/gt-foundation-2026-10-21.yaml"),
    "fusion_v2": ("artifacts/experiments/gt-primary-2026-10-11-v2/metrics.json", "Previous serving experiment (v2, dynamic-only)",
                  "/created_at", "configs/experiments/gt-primary-2026-10-11-v1.yaml"),
    "inventory": ("reports/dataset_inventory.json", "Dataset import inventory", "/generated_at", "data import"),
    "sources_yaml": ("configs/sources.yaml", "Dataset source manifest (DOIs, licences)", None, "source manifest"),
    "submission": ("configs/submission.yaml", "Team and submission metadata", None, "submission metadata"),
    "experiment_yaml": ("configs/experiment.yaml", "Event, availability and correction settings", None, "experiment config"),
    "manifest": ("reports/evidence_manifest.json", "Selected-claim manifest (23 groups)", None,
                 "scripts/evidence_manifest.py"),
    "warning_utility_md": ("reports/warning_utility.md", "Warning notice report", None, "scripts/warning_utility.py"),
    "data_card": ("docs/data_card.md", "Data card", None, "documentation"),
    "data_audit": ("reports/data_audit.md", "Data audit (measured)", None, "data audit"),
    "decisions": ("docs/decisions.md", "Decision log D-001…D-054", None, "documentation"),
    "ledger": (f"{VER}/ledger.json", "Final local verification ledger (f073a10)", "/created/utc",
               "scripts/verification_ledger.py"),
    "ledger_overnight": ("reports/verification/overnight/ledger.json", "Earlier verification ledger (history)",
                         "/created/utc", "scripts/verification_ledger.py"),
    "junit_pytest": (f"{VER}/junit/pytest.xml", "Saved pytest JUnit (f073a10)", None, "pytest"),
    "junit_e2e": (f"{VER}/junit/e2e.xml", "Saved Playwright JUnit, normal + recorded journeys (f073a10)", None,
                  "playwright"),
    "junit_public": (f"{VER}/junit/e2e-public.xml", "Saved Playwright JUnit, public synthetic mode (f073a10)", None,
                     "playwright"),
    "junit_summary": (f"{VER}/junit/e2e-summary.xml", "Saved Playwright JUnit, summary/export (f073a10)", None,
                      "playwright"),
    "readme": ("README.md", "README (claims as stated)", None, "documentation"),
}


def file_row(fid: str) -> dict[str, Any]:
    rel, what, date_ptr, ident = FILES[fid]
    p = ROOT / rel
    date = None
    if date_ptr and rel.endswith(".json") and load(rel) is not None:
        date = resolve(load(rel), date_ptr)
    if not date and tracked(rel):
        date = last_commit(rel)[1]
    if not date and p.exists():
        date = dt.datetime.fromtimestamp(p.stat().st_mtime, dt.UTC).isoformat(timespec="seconds") + " (file time)"
    if ident.startswith("configs/experiments/"):
        h, d = first_commit(ident)
        ident = f"protocol {ident} (first commit {h} {d})" if h else f"protocol {ident} (not committed)"
    return {"id": fid, "path": rel, "what": what, "date": date or "unavailable", "sha256": sha256(p) or "unavailable",
            "identity": ident, "status": public_status(rel)}


# --------------------------------------------------------------------------- source values (stored once)

class Values:
    def __init__(self) -> None:
        self.rows: list[dict[str, Any]] = []
        self.by_id: dict[str, dict[str, Any]] = {}

    def add(self, vid: str, fid: str, locator: str, unit: str, desc: str, value: Any = "__read__") -> None:
        if value == "__read__":
            doc = load(FILES[fid][0]) if FILES[fid][0].endswith(".json") else None
            value = resolve(doc, locator) if doc is not None else None
        if isinstance(value, bool):
            value = "TRUE" if value else "FALSE"
        row = {"id": vid, "file": fid, "locator": locator, "unit": unit, "desc": desc, "value": value}
        assert vid not in self.by_id, vid
        self.rows.append(row)
        self.by_id[vid] = row

    def val(self, vid: str) -> Any:
        return self.by_id[vid]["value"]


V = Values()


def fusion_values() -> None:
    F = "fusion"
    fr = "/forecast_results[label={l}&personalization={p}&split=test&horizon={h}]/{k}"
    models = [("pers", "persistence:cgm-only-v1", "off", "Persistence (last reading)"),
              ("trend", "linear_trend:cgm-only-v1", "off", "Linear trend"),
              ("ridgep", "ridge:dyn-prof-v2", "off", "Ridge, dynamic + profile"),
              ("v2a", "lightgbm:dyn-v2", "adaptive", "LightGBM dynamic-only + correction (previous demo model, v2)"),
              ("lgbp", "lightgbm:dyn-prof-v2", "off", "LightGBM + profile, population"),
              ("serv", "lightgbm:dyn-prof-v2", "adaptive", "LightGBM + profile + correction (serving)"),
              ("cbpa", "catboost:dyn-prof-v2", "adaptive", "CatBoost + profile + correction (not selected)")]
    for key, label, pers, name in models:
        for h in (30, 60):
            for k, unit in (("mae", "mg/dL"), ("rmse", "mg/dL"), ("n_origins", "origins"), ("n_patients", "people")):
                if k in ("n_origins", "n_patients") and key not in ("serv",):
                    continue
                V.add(f"fus.{key}.{k}{h}", F, fr.format(l=label, p=pers, h=h, k=k), unit,
                      f"Serving test (final test segment): {name}, {h} min, {k} (macro over participants)" if k == "mae"
                      else f"Serving test: {name}, {h} min, {k}")
    cmp_ = "/comparisons[comparison={c}&horizon={h}]/{k}"
    for key, title in (("sp", "serving minus persistence"), ("adapt", "lightgbm:dyn-prof-v2: adaptive minus population"),
                       ("lgbcb", "lightgbm minus catboost (dyn-prof-v2, population)"),
                       ("prof", "lightgbm: profile minus dynamic-only (population)")):
        for h in (30, 60):
            V.add(f"fus.cmp.{key}.d{h}", F, cmp_.format(c=title, h=h, k="mean_diff_mae"), "mg/dL",
                  f"Serving test paired difference, {title}, {h} min (A minus B; negative = A lower error)")
            V.add(f"fus.cmp.{key}.lo{h}", F, cmp_.format(c=title, h=h, k="ci95") + "[0]", "mg/dL", "95 % CI lower (participant bootstrap)")
            V.add(f"fus.cmp.{key}.hi{h}", F, cmp_.format(c=title, h=h, k="ci95") + "[1]", "mg/dL", "95 % CI upper (participant bootstrap)")
    ev = "/events[model=lightgbm:dyn-prof-v2 (adaptive)&split=test&horizon=60]/{k}"
    for k, unit, d in (("true_onsets", "onsets", "new sustained >180 onsets (eligible)"),
                       ("detected_onsets", "onsets", "onsets matched to a warning"),
                       ("warning_episodes", "episodes", "warning episodes"),
                       ("false_alarm_episodes", "episodes", "false-alarm episodes"),
                       ("precision", "ratio", "episode precision"),
                       ("false_alarms_per_observed_patient_day", "per observed patient-day", "saved false-alarm rate"),
                       ("observed_patient_days", "patient-days", "observed patient-day denominator (saved event metadata)"),
                       ("eligible_patient_days", "patient-days", "eligible patient-day denominator (saved event metadata)"),
                       ("false_alarms_per_patient_day", "per eligible patient-day", "saved false alarms per eligible patient-day"),
                       ("qualified_lead_median", "min", "median qualified lead among detected onsets"),
                       ("qualified_lead_n", "onsets", "onsets with a qualified lead (denominator of the median)"),
                       ("threshold", "mg/dL", "warning margin (score = max forecast − 180)"),
                       ("time_base", "text", "time-base definition of the day denominators")):
        V.add(f"fus.ev.{k}", F, ev.format(k=k), unit, f"Serving 60-min warning, final test: {d}")
    V.add("fus.ev.iqr_lo", F, ev.format(k="qualified_lead_iqr") + "[0]", "min", "Qualified lead IQR lower (detected onsets)")
    V.add("fus.ev.iqr_hi", F, ev.format(k="qualified_lead_iqr") + "[1]", "min", "Qualified lead IQR upper (detected onsets)")
    pev = "/events[model=persistence&split=test&horizon=60]/{k}"
    for k in ("detected_onsets", "false_alarm_episodes", "false_alarms_per_observed_patient_day"):
        V.add(f"fus.pev.{k}", F, pev.format(k=k), "count" if "rate" not in k else "", f"Persistence 60-min warning (fires at 150 mg/dL), final test: {k}")
    iv = "/intervals[label=lightgbm:dyn-prof-v2&personalization=adaptive&split=test&horizon={h}]/{k}"
    for h in (30, 60):
        for k, unit in (("coverage", "share"), ("width", "mg/dL"), ("q", "mg/dL"), ("per_patient_min", "share"),
                        ("per_patient_max", "share"), ("n", "origins")):
            V.add(f"fus.iv.{k}{h}", F, iv.format(h=h, k=k), unit, f"Served 80 % range, final test, {h} min: {k}")
    for seg in ("training", "validation", "calibration", "test"):
        for k, unit in (("origins", "origins"), ("issued_forecast_origins", "origins"), ("covered_hours", "hours"),
                        ("matched_targets_30", "targets"), ("matched_targets_60", "targets"),
                        ("distinct_new_onsets_60", "onsets")):
            V.add(f"fus.seg.{seg}.{k}", F, f"/cohort/by_segment/{seg}/{k}", unit, f"Cohort, {seg} segment: {k}")
    V.add("fus.participants", F, "/cohort/participants", "people", "Participants in the serving experiment (Type 2 group)")
    V.add("fus.excluded", F, "/excluded", "people", "CGMacros participants outside the Type 2 group (length of /excluded)",
          len(resolve(load(FILES[F][0]), "/excluded") or []) if load(FILES[F][0]) else None)
    V.add("fus.sens.meal60", F, "/sensitivity[experiment=optimistic meal logging&split=test&horizon=60]/mae", "mg/dL",
          "Optimistic meal-logging sensitivity (not eligible for serving), test, 60 min")
    V.add("fus.model_version", F, "/model_version", "text", "Model version evaluated")
    V.add("fus.created_at", F, "/created_at", "UTC", "Experiment run time")
    V.add("fus.selection_rule", F, "/selection/rule", "text", "Declared selection rule")
    # people better than persistence: aggregate count derived at build from per-participant MAE (no rows exported)
    m = load(FILES[F][0])
    for h in (30, 60):
        cnt = None
        if m:
            pp = {(r["patient_id"], r["label"], r["personalization"]): r["mae"] for r in m["per_patient_test"] if r["horizon"] == h}
            ids = {k[0] for k in pp}
            cnt = sum(1 for i in ids if pp[(i, "lightgbm:dyn-prof-v2", "adaptive")] < pp[(i, "persistence:cgm-only-v1", "off")])
        V.add(f"fus.better{h}", F, f"/per_patient_test (count of people with serving MAE < persistence MAE, {h} min)",
              "people", f"Serving test: participants with lower serving error than persistence, {h} min (derived count)", cnt)


def warning_utility_values() -> None:
    """Notice shares from saved predictions via scripts/warning_utility.py (re-checked there against saved metrics)."""
    rows: dict[str, Any] = {}
    if (ROOT / FILES["fusion_pred"][0]).exists():
        sys.path.insert(0, str(ROOT / "scripts"))
        import numpy as np  # noqa: PLC0415
        import warning_utility as WU  # noqa: PLC0415

        r = next(x for x in WU.WT.demo() if x["key"] == "demo")
        leads = np.array(r.get("_leads", []), float)
        rows = {"onsets": r["onsets"], "detected": r["detected"], "ge15": int((leads >= 15).sum()),
                "ge30": int((leads >= 30).sum()), "ge60": int((leads >= 60).sum()), "firing": r["firing"],
                "eligible": r["eligible"], "share_eligible": r["share_eligible"], "lead_median": r["lead_median"],
                "lead_q25": float(np.percentile(leads, 25)) if len(leads) else None,
                "lead_q75": float(np.percentile(leads, 75)) if len(leads) else None, "n_leads": int(len(leads))}
    loc = "scripts/warning_utility.py demo row (WT.demo, key=demo) → "
    for k, unit, d in (("onsets", "onsets", "eligible onsets (denominator of every notice share)"),
                       ("detected", "onsets", "detected onsets"),
                       ("ge15", "onsets", "onsets warned ≥ 15 min ahead (qualified lead)"),
                       ("ge30", "onsets", "onsets warned ≥ 30 min ahead (qualified lead)"),
                       ("ge60", "onsets", "onsets warned ≥ 60 min ahead (qualified lead)"),
                       ("firing", "origins", "eligible origins at which the warning was active"),
                       ("eligible", "origins", "eligible origins (time-in-warning denominator)"),
                       ("share_eligible", "share", "share of eligible origins in warning"),
                       ("lead_median", "min", "median qualified lead (detected onsets)"),
                       ("n_leads", "onsets", "number of qualified leads (detected onsets)"),
                       ("lead_q25", "min", "qualified lead 25th percentile"), ("lead_q75", "min", "qualified lead 75th percentile")):
        V.add(f"wu.{k}", "fusion_pred", loc + k, unit, f"Serving 60-min warning, final test: {d}", rows.get(k))


def trend_values() -> None:
    for h in (30, 60):
        for t in ("all", "falling_fast", "falling", "steady", "rising", "rising_fast"):
            base = f"/rows[horizon={h}&trend={t}]"
            V.add(f"tr.{t}.mae{h}", "trend", base + "/mae_macro", "mg/dL", f"Validation segment, {t}, {h} min: macro MAE")
            V.add(f"tr.{t}.cov{h}", "trend", base + "/range_coverage", "share", f"Validation segment, {t}, {h} min: served range coverage")
            V.add(f"tr.{t}.n{h}", "trend", base + "/origins", "origins", f"Validation segment, {t}, {h} min: origins")
        V.add(f"tr.bias_high{h}", "trend", f"/rows[horizon={h}&trend=all]/mean_signed_error_target_above_180", "mg/dL",
              f"Validation segment, {h} min: mean signed error (forecast − actual) when actual > 180")


def other_values() -> None:
    for fid in ("parity", "parity_final"):
        for k, unit in (("compared", "forecasts"), ("offline_forecasts", "forecasts"), ("missing_online", "forecasts"),
                        ("over_tolerance", "forecasts"), ("tolerance_mg_dl", "mg/dL"), ("max_abs_diff_mg_dl", "mg/dL"),
                        ("participants", "people"), ("outcome", "text"), ("model_version", "text")):
            V.add(f"{fid}.{k}", fid, f"/{k}", unit, f"{FILES[fid][1]}: {k}")
    n = "nested"
    mm = "/tracks/multimodal_5min/pipelines[pipeline={p}&mode={m}]/horizons/{h}/{k}"
    for key, p, md in (("sel", "selected", "final"), ("pers", "persistence:-", "off"), ("v2", "v2_recipe:dyn-v2", "v2_policy")):
        for h in (30, 60):
            V.add(f"nest.{key}.mae{h}", n, mm.format(p=p, m=md, h=h, k="mae"), "mg/dL", f"Nested internal (14 people held out once): {p} {md}, {h} min MAE")
    for h in (30, 60):
        V.add(f"nest.cov{h}", n, f"/tracks/multimodal_5min/pipelines[pipeline=selected&mode=final]/intervals/{h}/global/coverage_pooled",
              "share", f"Nested internal: selected procedure 80 % range coverage, {h} min")
        c = f"/tracks/multimodal_5min/comparisons[comparison=Profile inputs: LightGBM L2 dyn-prof-v2 minus dyn-v2&horizon={h}]"
        V.add(f"nest.prof.d{h}", n, c + "/mean_diff", "mg/dL", f"Nested: profile minus no profile (LightGBM L2, population), {h} min")
        V.add(f"nest.prof.lo{h}", n, c + "/ci95[0]", "mg/dL", "95 % CI lower")
        V.add(f"nest.prof.hi{h}", n, c + "/ci95[1]", "mg/dL", "95 % CI upper")
        ext = "/external/pipelines[pipeline={p}&mode={m}]/horizons/{h}/mae"
        V.add(f"ext.cand.mae{h}", n, ext.format(p="lightgbm_l2:cgm15-v1", m="final", h=h), "mg/dL", f"External ShanghaiT2DM (98): 15-min LightGBM, {h} min MAE")
        V.add(f"ext.pers.mae{h}", n, ext.format(p="persistence:-", m="off", h=h), "mg/dL", f"External ShanghaiT2DM: persistence, {h} min MAE")
        ec = f"/external/comparisons[comparison=Selected candidate minus persistence&horizon={h}]"
        for k in ("mean_diff", "ci95[0]", "ci95[1]", "improved", "n_participants"):
            V.add(f"ext.cmp.{k.replace('ci95[0]', 'lo').replace('ci95[1]', 'hi')}{h}", n, ec + "/" + k, "mg/dL" if "ci" in k or "mean" in k else "people",
                  f"External: candidate minus persistence, {h} min, {k}")
        V.add(f"ext.cov{h}", n, f"/external/pipelines[pipeline=lightgbm_l2:cgm15-v1&mode=final]/intervals/{h}/global/coverage_pooled",
              "share", f"External: 80 % range coverage, {h} min")
    V.add("nest.ridge_folds", n, "/tracks/multimodal_5min/selection_by_fold (count of outer folds selecting ridge:cgm-only-v1)",
          "folds", "Nested: outer folds whose rule chose Ridge CGM-only",
          sum(1 for k, v in (resolve(load(FILES[n][0]), "/tracks/multimodal_5min/selection_by_fold") or {}).items()
              if k.startswith("outer") and v.get("selected") == "ridge:cgm-only-v1") if load(FILES[n][0]) else None)
    V.add("nest.frozen_at", n, "/frozen_at", "UTC", "Nested protocol frozen")
    V.add("nest.opened_at", n, "/assessment_opened/opened_at", "UTC", "Nested internal assessment opened")
    t = "twoh"
    for h in (30, 60, 120):
        V.add(f"2h.sel.mae{h}", t, f"/forecasts/selected|final/horizons/{h}/mae", "mg/dL", f"Two-hour study: selected personalised, {h} min MAE")
        V.add(f"2h.pers.mae{h}", t, f"/forecasts/persistence:-|off/horizons/{h}/mae", "mg/dL", f"Two-hour study: persistence, {h} min MAE")
    c = "/comparisons[comparison=Selected (personalised) minus persistence&horizon=120]"
    V.add("2h.cmp.d", t, c + "/mean_difference", "mg/dL", "Two-hour: selected minus persistence at 120 min")
    V.add("2h.cmp.lo", t, c + "/ci95[0]", "mg/dL", "95 % CI lower")
    V.add("2h.cmp.hi", t, c + "/ci95[1]", "mg/dL", "95 % CI upper")
    V.add("2h.cmp.better", t, c + "/better_for", "people", "People with lower error than persistence")
    V.add("2h.cmp.n", t, c + "/participants", "people", "Participants")
    c = "/comparisons[comparison=Profile (EHR-like) inputs: LightGBM dyn-prof-v2 minus dyn-v2&horizon=120]"
    V.add("2h.prof.d", t, c + "/mean_difference", "mg/dL", "Two-hour profile ablation (LightGBM), 120 min")
    V.add("2h.prof.lo", t, c + "/ci95[0]", "mg/dL", "95 % CI lower")
    V.add("2h.prof.hi", t, c + "/ci95[1]", "mg/dL", "95 % CI upper")
    w = "/warnings[pipeline=selected&mode=final&horizon=120]"
    for k, unit in (("true_onsets", "onsets"), ("detected_onsets", "onsets"), ("false_alarms_per_observed_day", "per observed day"),
                    ("qualified_lead_median", "min"), ("qualified_lead_iqr[0]", "min"), ("qualified_lead_iqr[1]", "min")):
        V.add(f"2h.w.{k.replace('[0]', '_lo').replace('[1]', '_hi')}", t, f"{w}/{k}", unit, f"Two-hour 120-min warning: {k}")
    V.add("2h.cov120", t, "/forecasts/selected|final/coverage/120", "share", "Two-hour: 80 % range coverage at 120 min")
    V.add("2h.frozen_at", t, "/frozen_at", "UTC", "Two-hour protocol frozen")
    V.add("2h.opened_at", t, "/assessment_opened/opened_at", "UTC", "Two-hour assessment opened")
    e = "twoh_ext"
    for h in (30, 60, 120):
        V.add(f"2hx.mae{h}", e, f"/mae/{h}", "mg/dL", f"External two-hour (98 ShanghaiT2DM): {h} min MAE")
        V.add(f"2hx.pers{h}", e, f"/persistence_mae/{h}", "mg/dL", f"External two-hour: persistence {h} min MAE")
    for k in ("mean_diff", "ci95[0]", "ci95[1]", "improved", "n"):
        V.add(f"2hx.cmp.{k.replace('ci95[0]', 'lo').replace('ci95[1]', 'hi')}", e, f"/minus_persistence/120/{k}", "mg/dL" if "ci" in k or "mean" in k else "people",
              f"External two-hour: minus persistence at 120 min, {k}")
    V.add("2hx.cov120", e, "/coverage/120", "share", "External two-hour: 80 % range coverage at 120 min")
    V.add("2hx.participants", e, "/participants", "people", "External two-hour participants")
    V.add("2hx.opened_at", e, "/opened/opened_at", "UTC", "External two-hour assessment opened")
    c = "clf"
    for k, unit in (("detected_onsets", "onsets"), ("true_onsets", "onsets"), ("false_alarms_per_observed_day", "per observed day"),
                    ("share_of_eligible_time_warning", "share"), ("origin_roc_auc", "AUC")):
        V.add(f"clf.{k}", c, f"/results/selected/{k}", unit, f"Two-hour classifier (D-044): {k}")
    V.add("clf.opened_at", c, "/opened/opened_at", "UTC", "Classifier assessment opened")
    V.add("clf.reopened_at", c, "/opened/reopened[0]/at", "UTC", "Classifier assessment re-opened (post-hoc diagnostic)")
    r = "warnrange"
    for arm, key in (("B0_calibration", "b0"), ("B1_pooled", "b1")):
        for h in (60, 120):
            V.add(f"wr.{key}.cov{h}", r, f"/ranges/{arm}/{h}/coverage_pooled", "share", f"D-046 ranges {arm}, {h} min coverage")
            V.add(f"wr.{key}.width{h}", r, f"/ranges/{arm}/{h}/width_mean", "mg/dL", f"D-046 ranges {arm}, {h} min mean width")
    V.add("wr.b1.pmed120", r, "/ranges/B1_pooled/120/coverage_participant_median", "share", "D-046 pooled ranges: per-person median coverage, 120 min")
    V.add("wr.early.sel", r, "/warnings/selected/early", "onsets", "D-046 A: onsets warned ≥ 30 min ahead, selected arm")
    V.add("wr.early.inc", r, "/warnings/incumbent/early", "onsets", "D-046 A: onsets warned ≥ 30 min ahead, incumbent")
    V.add("wr.onsets", r, "/warnings/selected/onsets", "onsets", "D-046 A: onsets (denominator)")
    V.add("wr.lead.sel", r, "/warnings/selected/lead_median", "min", "D-046 A: median lead, selected arm")
    V.add("wr.diff.d", r, "/warning_diff_selected_minus_incumbent/early_share_diff", "share", "D-046 A: early-share difference selected − incumbent")
    V.add("wr.diff.lo", r, "/warning_diff_selected_minus_incumbent/early_share_ci95[0]", "share", "95 % CI lower")
    V.add("wr.diff.hi", r, "/warning_diff_selected_minus_incumbent/early_share_ci95[1]", "share", "95 % CI upper")
    for seg, key in (("all", "all"), ("first_24h", "d1")):
        p = f"/profile_prior/{seg}/120/c1_minus_population"
        V.add(f"wr.prior.{key}.d", r, p + "/mean_diff", "mg/dL", f"D-046 C: profile prior minus population, 120 min ({seg})")
        V.add(f"wr.prior.{key}.lo", r, p + "/ci95[0]", "mg/dL", "95 % CI lower")
        V.add(f"wr.prior.{key}.hi", r, p + "/ci95[1]", "mg/dL", "95 % CI upper")
    V.add("wr.opened_at", r, "/opened/opened_at", "UTC", "D-046 outer predictions read")
    b = "blend"
    for k in ("improvement_60", "worsening_30", "high_worsening_60", "availability_drop_points"):
        V.add(f"blend.{k}", b, f"/decision/checks/{k}", "mg/dL" if "points" not in k else "points", f"D-048 development check: {k} (improvement positive = gain)")
    V.add("blend.alpha", b, "/decision/final_selected_alpha", "alpha", "D-048 selected blend weight")
    V.add("blend.target_met", b, "/decision/target_met", "flag", "D-048 saved outcome: development target met")
    V.add("blend.inc60", b, "/folds/final/alphas/1.0/mae_60", "mg/dL", "D-048 alpha = 1 (incumbent) 60-min development MAE")
    V.add("blend.written_at", "blend_proto", "/written_at", "UTC", "D-048 protocol written (before any alpha < 1 scored)")
    bp = ROOT / FILES[b][0]
    V.add("blend.dev_time", b, "file modification time of development.json (local; the file has no timestamp field)", "UTC",
          "D-048 development decision written", dt.datetime.fromtimestamp(bp.stat().st_mtime, dt.UTC).replace(tzinfo=None, microsecond=0)
          if bp.exists() else None)
    t = "retrain"
    for arm in ("A1_participant_balanced", "A2_balanced_rising_2x", "B1_balanced_long_history"):
        a = arm.split("_")[0]
        for k in ("gain_60", "worsening_30", "high_worsening_60", "availability_drop_points"):
            V.add(f"rt.{a}.{k}", t, f"/selection/by_arm/{arm}/checks/{k}", "mg/dL" if "points" not in k else "points",
                  f"D-049 {arm} vs B0: {k} (gain positive = lower error)")
        V.add(f"rt.{a}.qualifies", t, f"/selection/by_arm/{arm}/qualifies", "flag", f"D-049 {arm}: saved qualifies flag")
    V.add("rt.B1pop.d", t, "/development_paired_descriptive/B1_pop_60/mean_diff", "mg/dL", "D-049 B1 population-only 60-min change (descriptive, after selection)")
    V.add("rt.B1pop.lo", t, "/development_paired_descriptive/B1_pop_60/ci95[0]", "mg/dL", "95 % CI lower")
    V.add("rt.B1pop.hi", t, "/development_paired_descriptive/B1_pop_60/ci95[1]", "mg/dL", "95 % CI upper")
    V.add("rt.started_at", t, "/started_at", "UTC", "D-049 run start")
    g = "ranges"
    for s in ("G", "T"):
        V.add(f"rg.err.{s}", g, f"/pre_test_check/coverage_error/{s}", "share", f"D-050 pre-test mean |coverage − 0.80|, shape {s}")
        V.add(f"rg.width.{s}", g, f"/pre_test_check/width_mean/{s}", "mg/dL", f"D-050 pre-test mean width, shape {s}")
    V.add("rg.chosen", g, "/pre_test_check/chosen", "text", "D-050 saved choice")
    for arm, key in (("incumbent_G_calibration", "inc"), ("G_pooled", "pool")):
        for h in (30, 60):
            V.add(f"rg.{key}.cov{h}", g, f"/test_retrospective/{arm}/{h}/coverage", "share", f"D-050 final test (read once), {arm}, {h} min coverage")
            V.add(f"rg.{key}.width{h}", g, f"/test_retrospective/{arm}/{h}/width_mean", "mg/dL", f"D-050 final test, {arm}, {h} min mean width")
            for st in ("rising", "steady", "falling"):
                V.add(f"rg.{key}.{st}{h}", g, f"/test_retrospective/{arm}/{h}/by_stratum/{st}/coverage", "share",
                      f"D-050 final test, {arm}, {h} min, {st} (> 1 mg/dL/min) coverage")
                if key == "inc":
                    V.add(f"rg.{key}.{st}n{h}", g, f"/test_retrospective/{arm}/{h}/by_stratum/{st}/n", "origins", f"D-050 final test, {st}, {h} min origins")
    V.add("rg.generated_at", g, "/generated_at", "UTC", "D-050 run time")
    u = "tune"
    for arm, key in (("C0_current_features", "C0"), ("C1_with_history", "C1"), ("R0_untuned_serving_recipe", "R0"), ("persistence", "P")):
        V.add(f"tu.{key}.mae60", u, f"/arms/{arm}/60/mae", "mg/dL", f"D-051 {arm}: 60-min personalised macro MAE")
    for pair, key in (("C1_with_history|C0_current_features|final|60", "c1c0"), ("C1_with_history|R0_untuned_serving_recipe|final|60", "c1r0")):
        for k in ("mean_diff", "ci95[0]", "ci95[1]", "improved"):
            V.add(f"tu.{key}.{k.replace('ci95[0]', 'lo').replace('ci95[1]', 'hi')}", u, f"/paired/{pair}/{k}", "mg/dL" if k != "improved" else "people",
                  f"D-051 paired {pair}: {k}")
    for k in ("worsening_30", "high_worsening_60", "c1_better"):
        V.add(f"tu.adopt.{k}", u, f"/adoption/{k}", "mg/dL" if k != "c1_better" else "flag", f"D-051 adoption check: {k}")
    V.add("tu.opened_at", u, "/opened/opened_at", "UTC", "D-051 outer folds opened")
    V.add("tu.started_at", u, "/started_at", "UTC", "D-051 run start")
    j = "joint"
    for p_ in ("S", "R0"):
        V.add(f"jt.{p_}.mae60", j, f"/procedures/{p_}/60/mae", "mg/dL", f"D-052 {p_}: 60-min personalised macro MAE")
        V.add(f"jt.{p_}.share30", j, f"/warnings/{p_}/share_ge30", "share", f"D-052 {p_}: share of all onsets warned ≥ 30 min")
        V.add(f"jt.{p_}.fa", j, f"/warnings/{p_}/fa_per_observed_day", "per observed day", f"D-052 {p_}: false alarms per observed day")
    for k in ("mean_diff_60", "ci95_60[0]", "ci95_60[1]", "worsening_30", "high_worsening_60", "availability_drop_points", "passed"):
        V.add(f"jt.gate.{k.replace('[0]', '_lo').replace('[1]', '_hi')}", j, f"/gate/{k}", "flag" if k == "passed" else "mg/dL",
              f"D-052 primary gate value: {k}")
    for k in ("fa_per_day_delta", "time_in_warning_delta_points", "warned30_delta_points", "passed"):
        V.add(f"jt.rel.{k}", j, f"/release_guardrails/{k}", "flag" if k == "passed" else "", f"D-052 release guardrail value: {k}")
    for h in (30, 60):
        V.add(f"jt.cov{h}", j, f"/ranges/S/{h}/coverage", "share", f"D-052 candidate S range coverage, {h} min")
    V.add("jt.started_at", j, "/started_at", "UTC", "D-052 run start")
    V.add("jt.opened_at", j, "/opened/opened_at", "UTC", "D-052 outer folds opened")
    c = "constrained"
    for p_ in ("S", "R0", "persistence"):
        for h in (30, 60):
            V.add(f"cs.{p_}.mae{h}", c, f"/procedures/{p_}/{h}/mae", "mg/dL", f"D-053 {p_}: {h}-min personalised macro MAE")
    for p_ in ("S", "R0"):
        V.add(f"cs.{p_}.high60", c, f"/procedures/{p_}/60/mae_high", "mg/dL", f"D-053 {p_}: 60-min macro MAE on targets > 180")
    for k in ("mean_diff", "ci95[0]", "ci95[1]", "improved"):
        V.add(f"cs.p60.{k.replace('ci95[0]', 'lo').replace('ci95[1]', 'hi')}", c, f"/paired/S|R0|final|60/{k}", "mg/dL" if k != "improved" else "people",
              f"D-053 paired S − R0 at 60 min: {k}")
    for k in ("mean_diff_60", "ci95_60[1]", "worsening_30", "high_worsening_60", "availability_drop_points"):
        V.add(f"cs.pg.{k.replace('[1]', '_hi')}", c, f"/primary_gate/values/{k}", "mg/dL" if "points" not in k else "points",
              f"D-053 primary gate value (unrounded): {k}")
        V.add(f"cs.pgok.{k.replace('ci95_60[1]', 'ci95_upper_below_0').replace('availability_drop_points', 'availability_drop')}", c,
              f"/primary_gate/checks/{k.replace('ci95_60[1]', 'ci95_upper_below_0').replace('availability_drop_points', 'availability_drop')}",
              "flag", f"D-053 saved check outcome: {k}")
    V.add("cs.pg.passed", c, "/primary_gate/passed", "flag", "D-053 saved primary gate outcome")
    for p_, key in (("S", "S"), ("R0-P", "R0P")):
        for k, unit in (("fa_day", "per observed day"), ("share30", "share"), ("tiw", "share"), ("false_alarms", "episodes"),
                        ("observed_days", "patient-days"), ("onsets", "onsets")):
            V.add(f"cs.w.{key}.{k}", c, f"/warnings/{p_}/{k}", unit, f"D-053 warning {p_}: {k}")
    V.add("cs.op.passed", c, "/operational_gate/passed", "flag", "D-053 saved operational gate outcome")
    V.add("cs.rg.passed", c, "/range_gate/passed", "flag", "D-053 saved range gate outcome")
    V.add("cs.all_passed", c, "/all_gates_passed", "flag", "D-053 saved all-gates outcome")
    for h in (30, 60):
        V.add(f"cs.cov{h}", c, f"/ranges/S/{h}/coverage", "share", f"D-053 S exact-predictor 80 % range coverage, {h} min")
    V.add("cs.started_at", c, "/manifest/started_at", "UTC", "D-053 run start")
    V.add("cs.opened_at", c, "/assessment_opened/opened_at", "UTC", "D-053 assessment opened")
    V.add("cs.fits", c, "/fit_totals/fits", "fits", "D-053 logical fits")
    f = "foundation"
    V.add("fd.FM60", f, "/horizons/60/FM_vs_persistence/macro/FM", "mg/dL", "D-054 zero-shot foundation model, 60-min macro MAE")
    V.add("fd.FM30", f, "/horizons/30/FM_vs_persistence/macro/FM", "mg/dL", "D-054 foundation model, 30-min macro MAE")
    V.add("fd.P60", f, "/horizons/60/FM_vs_persistence/macro/persistence", "mg/dL", "D-054 persistence on identical origins, 60 min")
    V.add("fd.P30", f, "/horizons/30/FM_vs_persistence/macro/persistence", "mg/dL", "D-054 persistence on identical origins, 30 min")
    V.add("fd.R060", f, "/horizons/60/FM_vs_R0saved_final/macro/R0", "mg/dL", "D-054 refitted serving recipe (saved D-052 R0), 60 min")
    for key, blk in (("fp", "FM_vs_persistence"), ("fr", "FM_vs_R0saved_final"), ("f1", "F1_final_vs_R0refit_final")):
        for k in ("mean", "lo", "hi", "people_better"):
            V.add(f"fd.{key}.{k}", f, f"/horizons/60/{blk}/diff_macro/{k}", "mg/dL" if k != "people_better" else "people",
                  f"D-054 {blk}, 60 min: {k}")
    V.add("fd.high.FM", f, "/horizons/60/FM_vs_R0saved_final/high_gt180/macro/FM", "mg/dL", "D-054 foundation, targets > 180, 60 min")
    V.add("fd.high.R0", f, "/horizons/60/FM_vs_R0saved_final/high_gt180/macro/R0", "mg/dL", "D-054 R0, targets > 180, 60 min")
    V.add("fd.generated_at", f, "/generated_at", "UTC", "D-054 summary generated")
    V.add("fd.checkpoint", "foundation_meta", "/checkpoint", "text", "D-054 checkpoint")
    V.add("fd.revision", "foundation_meta", "/revision", "text", "D-054 pinned checkpoint revision")
    V.add("fd.license", "foundation_meta", "/license", "text", "D-054 checkpoint licence")
    V.add("fd.origins", "foundation_meta", "/origins", "origins", "D-054 origins scored")
    V.add("inv.cgm", "inventory", "/datasets[dataset_id=cgmacros]/parsed_participants", "people", "CGMacros participants imported")
    V.add("inv.sh", "inventory", "/datasets[dataset_id=shanghai_t2dm]/parsed_participants", "people", "ShanghaiT2DM participants imported")
    V.add("inv.cgm_eval", "inventory", "/assessed_participants_by_dataset/cgmacros", "people", "CGMacros participants evaluated")
    V.add("inv.sh_eval", "inventory", "/assessed_participants_by_dataset/shanghai_t2dm", "people", "ShanghaiT2DM participants evaluated")
    V.add("reg.serving", "registry", "/serving/dataset_replay", "text", "Registry serving model, recorded data")
    V.add("reg.synthetic", "registry", "/serving/synthetic_demo", "text", "Registry serving model, synthetic mode")
    V.add("reg.synthetic_fitted", "registry", "/models[model_version=synthetic-demo.lightgbm.dyn-v2.v1]/synthetic_fitted",
          "flag", "Registry: synthetic demo model fitted on fictional data")
    led = load(FILES["ledger"][0]) or {"entries": []}
    for key, suite in (("pytest", "backend pytest (complete)"), ("vitest", "frontend unit (vitest)"),
                       ("e2e", "e2e normal + recorded judge journeys"), ("e2e_public", "e2e public mode (local, synthetic)"),
                       ("e2e_summary", "e2e summary/export")):
        idx = next((i for i, e in enumerate(led["entries"]) if e.get("suite") == suite), None)
        for k in ("passed", "skipped", "failed"):
            V.add(f"ver.{key}.{k}", "ledger", f"/entries[{idx}]/counts/{k}", "tests", f"Final ledger (f073a10): {suite}, {k}",
                  resolve(led, f"/entries[{idx}]/counts/{k}") if idx is not None else None)
    V.add("ver.commit", "ledger", "/snapshots[label=before]/revision/commit", "text", "Tested commit (final ledger)")
    # dataset / licence facts (text, from the tracked source manifest)
    import yaml  # noqa: PLC0415

    src = yaml.safe_load((ROOT / FILES["sources_yaml"][0]).read_text())
    ds = (src or {}).get("sources") or {}
    for key, did in (("cgm", "cgmacros"), ("sh", "shanghai_t2dm")):
        d = ds.get(did, {})
        V.add(f"src.{key}.doi", "sources_yaml", f"sources.{did}.doi", "text", f"{did} DOI", d.get("doi"))
        V.add(f"src.{key}.license", "sources_yaml", f"sources.{did}.license", "text", f"{did} licence", d.get("license"))
        V.add(f"src.{key}.version", "sources_yaml", f"sources.{did}.release", "text", f"{did} release", d.get("release"))
    V.add("src.cgm.publisher_n", "sources_yaml", "sources.cgmacros.publisher_participants", "people",
          "CGMacros participants per the publisher", ds.get("cgmacros", {}).get("publisher_participants"))
    dex = next((x for x in ds.get("cgmacros", {}).get("devices", []) if x.get("device_id") == "dexcom_g6_pro"), {})
    cfg = yaml.safe_load((ROOT / FILES["experiment_yaml"][0]).read_text())
    ev_cfg = _find_key(cfg, "confirm_minutes")
    V.add("cfg.threshold", "experiment_yaml", "events.threshold_mg_dl", "mg/dL", "Event threshold (strictly greater than)",
          (ev_cfg or {}).get("threshold_mg_dl"))
    V.add("cfg.confirm", "experiment_yaml", "events.confirm_minutes", "min", "Minutes above threshold to confirm an episode",
          (ev_cfg or {}).get("confirm_minutes"))
    V.add("cfg.stale", "experiment_yaml", "features.stale_threshold_minutes", "min", "Stale after this many minutes without a reading",
          (_find_key(cfg, "stale_threshold_minutes") or {}).get("stale_threshold_minutes"))
    V.add("cfg.lookback", "experiment_yaml", "features.lookback_minutes", "min", "Feature lookback window",
          (_find_key(cfg, "lookback_minutes") or {}).get("lookback_minutes"))
    V.add("cfg.reset_gap", "experiment_yaml", "adaptation.reset_gap_minutes", "min", "Correction resets after a gap of this length",
          (_find_key(cfg, "reset_gap_minutes") or {}).get("reset_gap_minutes"))
    V.add("const.min_per_day", "experiment_yaml", "constant: minutes per day", "min", "Minutes in a day (definition)", 1440)
    V.add("const.cadence", "sources_yaml", "sources.cgmacros.devices[dexcom_g6_pro].native_cadence_minutes", "min",
          "Dexcom G6 Pro native cadence", dex.get("native_cadence_minutes"))
    V.add("const.target_cov", "experiment_yaml", "intervals level 0.8 (declared target coverage)", "share",
          "Declared prediction-range target coverage", (_find_key(cfg, "level") or {}).get("level", 0.8))


SLIDE9 = ["lightgbm:dyn-prof-v2", "catboost:dyn-prof-v2", "lightgbm:dyn-v2", "catboost:dyn-v2", "ridge:dyn-v2",
          "ridge:dyn-prof-v2", "linear_trend:cgm-only-v1", "persistence:cgm-only-v1"]


def deck_values() -> None:
    """Values that appear only on slides 6–15 or in the narration."""
    for lab in SLIDE9:
        for k, d in (("selection_mae_30", "validation MAE 30"), ("selection_mae_60", "validation MAE 60"),
                     ("selection_score", "selection score (mean of 30/60)")):
            V.add(f"sel.{lab.replace(':', '_')}.{k}", "fusion", f"/selection/candidates[label={lab}]/{k}", "mg/dL",
                  f"Serving experiment selection (validation segment, population model): {lab}, {d}")
    V.add("sel.selected", "fusion", "/selection/selected_label", "text", "Selected candidate (fusion-constrained rule)")
    V.add("sel.v2_selected", "fusion_v2", "/selection/selected_label", "text", "Selected candidate without the fusion constraint (v2)")
    f = "foundation"
    for k in ("mean", "lo", "hi", "people_better"):
        V.add(f"fd.fp30.{k}", f, f"/horizons/30/FM_vs_persistence/diff_macro/{k}", "mg/dL" if k != "people_better" else "people",
              f"D-054 foundation minus persistence, 30 min: {k}")
    V.add("fd.params", "foundation_meta", "/parameters", "parameters", "D-054 checkpoint parameter count")


def _find_key(o: Any, key: str) -> dict | None:
    """First mapping in a nested YAML document that contains ``key``."""
    if isinstance(o, dict):
        if key in o:
            return o
        for v in o.values():
            r = _find_key(v, key)
            if r is not None:
                return r
    elif isinstance(o, list):
        for v in o:
            r = _find_key(v, key)
            if r is not None:
                return r
    return None


# --------------------------------------------------------------------------- claims
# A claim row: id, exact claim, locations, scope, declared status, expr, stated, decimals, limitation, action.
# ``expr`` is arithmetic over {value-id} placeholders (and the Excel/Python functions ABS, MIN, MAX); the workbook writes
# it as a formula over Sources cells and the builder evaluates the same expression in Python. A numeric claim whose value
# does not round to ``stated`` at ``dec`` decimals is set to unsupported/contradicted; a missing source sets pending.
# ``expr`` may be None for non-numeric claims (then ``sources`` names the evidence: value ids, file ids, or test refs).

def C(cid: str, text: str, locs: str, scope: str, status: str, expr: str | None = None, stated: float | str | None = None,
      dec: int | None = None, limitation: str = "", action: str = "", sources: str = "", manifest: str = "") -> dict:
    return dict(id=cid, text=text, locs=locs, scope=scope, status=status, expr=expr, stated=stated, dec=dec,
                limitation=limitation, action=action, sources=sources, manifest=manifest)


SERV = "Serving model, final test segment (later recordings of the same 14 people; retrospective, inspected)"
NEST = "Participant-held-out internal (14 CGMacros people, each held out once; previously inspected)"
EXT = "External ShanghaiT2DM, 98 people, CGM-only procedure (not the serving fusion model)"
RT = "README.md (Results at a glance)"


def manifest_claims() -> list[dict]:
    """The 23 evidence-manifest groups expanded to one row per checked number (manifest = 'group | check label')."""
    L = f"{RT}; reports/evidence_manifest.md"
    out = [
        C("C-SRV-01", "The app serves lightgbm.dyn-prof-v2.gt-primary-2026-10-13-fusion for recorded data", L + "; README §7",
          "Software configuration", "supported", sources="reg.serving; registry",
          manifest="serving-identity | serving model (recorded data)"),
        C("C-FC-01", "Serving MAE 12.96 mg/dL at 30 min", L, SERV, "supported", "{fus.serv.mae30}", 12.96, 2,
          "Same 14 people as training; macro over participants", manifest="demo-forecast | MAE 30"),
        C("C-FC-02", "Serving MAE 20.26 mg/dL at 60 min", L, SERV, "supported", "{fus.serv.mae60}", 20.26, 2,
          "Same 14 people as training", manifest="demo-forecast | MAE 60"),
        C("C-FC-03", "Persistence MAE 14.84 mg/dL at 30 min (same origins)", L, SERV, "supported", "{fus.pers.mae30}", 14.84, 2,
          manifest="demo-forecast | persistence 30"),
        C("C-FC-04", "Persistence MAE 23.23 mg/dL at 60 min (same origins)", L, SERV, "supported", "{fus.pers.mae60}", 23.23, 2,
          manifest="demo-forecast | persistence 60"),
        C("C-WN-01", "61 new high-glucose onsets warned (60-min warning)", L, SERV, "supported", "{fus.ev.detected_onsets}", 61, 0,
          "Research event definition; not triage", manifest="demo-warning | onsets warned"),
        C("C-WN-02", "64 new high-glucose onsets in the final test segment", L, SERV, "supported", "{fus.ev.true_onsets}", 64, 0,
          "Onsets ≠ origins ≠ participants", manifest="demo-warning | onsets"),
        C("C-WN-03", "40 false alarms", L, SERV, "supported", "{fus.ev.false_alarm_episodes}", 40, 0,
          manifest="demo-warning | false alarms"),
        C("C-WN-04", "1.49 false alarms per observed patient-day", L, SERV, "supported",
          "{fus.ev.false_alarm_episodes}/{fus.ev.observed_patient_days}", 1.49, 2,
          "Denominator = observed patient-days from saved event metadata (approximate time base: issued origins × 5 min)",
          manifest="demo-warning | false alarms per observed day"),
        C("C-WN-05", "Median forecast lead 30 min among detected onsets", L, SERV, "supported", "{fus.ev.qualified_lead_median}", 30, 0,
          "Denominator = detected onsets with a qualified lead (61), not all onsets", manifest="demo-warning | median forecast lead"),
        C("C-WN-06", "Warning active at 11.8 % of eligible times", L + "; reports/warning_time_share.md", SERV, "supported",
          "{wu.firing}/{wu.eligible}*100", 11.8, 1, "Post-hoc diagnostic", manifest="demo-time-in-warning | share of eligible times warning"),
        C("C-RG-01", "Served 80 % ranges covered 74.2 % at 30 min (under-coverage)", L, SERV, "supported", "{fus.iv.coverage30}*100", 74.2, 1,
          "Pooled over origins; per-person coverage varies", manifest="demo-ranges | demo coverage 30"),
        C("C-RG-02", "Served 80 % ranges covered 74.0 % at 60 min (under-coverage)", L, SERV, "supported", "{fus.iv.coverage60}*100", 74.0, 1,
          manifest="demo-ranges | demo coverage 60"),
        C("C-RG-03", "Research procedure ranges covered 80.1 % at 30 min (people held out)", L, NEST, "supported", "{nest.cov30}*100", 80.1, 1,
          "Different procedure from the served model", manifest="demo-ranges | nested research coverage 30"),
        C("C-RG-04", "Research procedure ranges covered 78.8 % at 60 min (people held out)", L, NEST, "supported", "{nest.cov60}*100", 78.8, 1,
          manifest="demo-ranges | nested research coverage 60"),
        C("C-NS-01", "Research candidate MAE 13.88 at 30 min (D-032)", L, NEST, "supported", "{nest.sel.mae30}", 13.88, 2,
          "Not served; cohort previously inspected", manifest="nested-forecast | MAE 30"),
        C("C-NS-02", "Research candidate MAE 22.30 at 60 min (D-032)", L, NEST, "supported", "{nest.sel.mae60}", 22.30, 2,
          manifest="nested-forecast | MAE 60"),
        C("C-NS-03", "Nested persistence 15.34 at 30 min", L, NEST, "supported", "{nest.pers.mae30}", 15.34, 2,
          manifest="nested-forecast | persistence 30"),
        C("C-NS-04", "Nested persistence 24.47 at 60 min", L, NEST, "supported", "{nest.pers.mae60}", 24.47, 2,
          manifest="nested-forecast | persistence 60"),
        C("C-NS-05", "Ridge on CGM only chosen in 4 of 5 folds", L, NEST, "supported", "{nest.ridge_folds}", 4, 0,
          manifest="nested-forecast | folds choosing Ridge CGM-only"),
        C("C-EX-01", "External 15-min LightGBM MAE 11.15 at 30 min", L, EXT, "supported", "{ext.cand.mae30}", 11.15, 2,
          "CGM only; tests the forecasting core, not the fusion", manifest="external-15min | MAE 30"),
        C("C-EX-02", "External 15-min LightGBM MAE 20.11 at 60 min", L, EXT, "supported", "{ext.cand.mae60}", 20.11, 2,
          manifest="external-15min | MAE 60"),
        C("C-EX-03", "External persistence 12.14 at 30 min", L, EXT, "supported", "{ext.pers.mae30}", 12.14, 2,
          manifest="external-15min | persistence 30"),
        C("C-EX-04", "External persistence 20.44 at 60 min", L, EXT, "supported", "{ext.pers.mae60}", 20.44, 2,
          manifest="external-15min | persistence 60"),
        C("C-EX-05", "External 60-minute difference is inconclusive (95 % CI upper bound above 0)", L, EXT, "supported",
          "{ext.cmp.hi60}", None, None, "Interval crosses zero", sources="ext.cmp.lo60; ext.cmp.hi60",
          manifest="external-15min | 60-minute difference"),
        C("C-PF-01", "Profile minus no profile (held out) +0.34 mg/dL at 30 min", L + "; README §7", NEST, "negative-result",
          "{nest.prof.d30}", 0.34, 2, "CI [−0.28, +0.89] crosses zero: no reliable benefit", manifest="profile-ablation-nested | profile minus no profile, 30 min"),
        C("C-PF-02", "Profile minus no profile (held out) −0.75 mg/dL at 60 min", L + "; README §7", NEST, "negative-result",
          "{nest.prof.d60}", -0.75, 2, "CI [−2.19, +0.61] crosses zero", manifest="profile-ablation-nested | profile minus no profile, 60 min"),
        C("C-PF-03", "Profile ablation 30-min CI lower −0.28", L, NEST, "supported", "{nest.prof.lo30}", -0.28, 2),
        C("C-PF-04", "Profile ablation 30-min CI upper +0.89", L, NEST, "supported", "{nest.prof.hi30}", 0.89, 2),
        C("C-PF-05", "Profile ablation 60-min CI lower −2.19", L, NEST, "supported", "{nest.prof.lo60}", -2.19, 2),
        C("C-PF-06", "Profile ablation 60-min CI upper +0.61", L, NEST, "supported", "{nest.prof.hi60}", 0.61, 2),
        C("C-2H-01", "Two hours ahead: MAE 30.6 at 120 min (D-043)", L, NEST, "supported", "{2h.sel.mae120}", 30.6, 1,
          "Research procedure, not served", manifest="two-hour | MAE 120"),
        C("C-2H-02", "Two-hour persistence 35.9 at 120 min", L, NEST, "supported", "{2h.pers.mae120}", 35.9, 1, manifest="two-hour | persistence 120"),
        C("C-2H-03", "Two-hour procedure better than persistence for 14 of 14 people", L, NEST, "supported", "{2h.cmp.better}", 14, 0,
          manifest="two-hour | people better"),
        C("C-2H-04", "Two-hour warning: 315 of 329 onsets", L, NEST, "supported", "{2h.w.detected_onsets}", 315, 0,
          manifest="two-hour | onsets warned"),
        C("C-2H-05", "Two-hour warning onsets (denominator) 329", L, NEST, "supported", "{2h.w.true_onsets}", 329, 0),
        C("C-2H-06", "Two-hour warning 1.62 false alarms per observed day", L, NEST, "supported", "{2h.w.false_alarms_per_observed_day}", 1.62, 2,
          manifest="two-hour | false alarms per observed day"),
        C("C-2H-07", "Two-hour warning median lead 15 min", L, NEST, "supported", "{2h.w.qualified_lead_median}", 15, 0,
          "Median among detected onsets", manifest="two-hour | median lead (IQR)"),
        C("C-2H-08", "Two-hour lead IQR 5–65 min (lower)", L, NEST, "supported", "{2h.w.qualified_lead_iqr_lo}", 5, 0),
        C("C-2H-09", "Two-hour lead IQR 5–65 min (upper)", L, NEST, "supported", "{2h.w.qualified_lead_iqr_hi}", 65, 0),
        C("C-2H-10", "Two-hour ranges covered 76.5 % (target 80 %)", L, NEST, "supported", "{2h.cov120}*100", 76.5, 1,
          manifest="two-hour | coverage 120"),
        C("C-2H-11", "Two-hour profile ablation −0.86 [−2.60, +1.02] (LightGBM)", "README §7; reports/two_hour_benchmark.md", NEST,
          "negative-result", "{2h.prof.d}", -0.86, 2, "CI crosses zero", manifest="two-hour | profile ablation 120 (LightGBM)"),
        C("C-2X-01", "External two-hour MAE 29.55 at 120 min (D-045)", L, EXT, "supported", "{2hx.mae120}", 29.55, 2,
          "Second opening of this cohort", manifest="two-hour-external | MAE 120"),
        C("C-2X-02", "External two-hour persistence 31.68", L, EXT, "supported", "{2hx.pers120}", 31.68, 2, manifest="two-hour-external | persistence 120"),
        C("C-2X-03", "External two-hour difference −2.12 mg/dL", L, EXT, "supported", "{2hx.cmp.mean_diff}", -2.12, 2,
          manifest="two-hour-external | difference (95 % CI)"),
        C("C-2X-04", "External two-hour CI lower −2.67", L, EXT, "supported", "{2hx.cmp.lo}", -2.67, 2),
        C("C-2X-05", "External two-hour CI upper −1.53", L, EXT, "supported", "{2hx.cmp.hi}", -1.53, 2),
        C("C-2X-06", "Better for 78 of 98 people (two-hour external)", L, EXT, "supported", "{2hx.cmp.improved}", 78, 0,
          manifest="two-hour-external | people better"),
        C("C-2X-07", "External two-hour 98 participants", L, EXT, "supported", "{2hx.cmp.n}", 98, 0),
        C("C-2X-08", "External two-hour MAE 10.90 at 30 min", L, EXT, "supported", "{2hx.mae30}", 10.90, 2, manifest="two-hour-external | MAE 30 / 60"),
        C("C-2X-09", "External two-hour MAE 19.31 at 60 min", L, EXT, "supported", "{2hx.mae60}", 19.31, 2),
        C("C-2X-10", "External two-hour persistence 12.14 at 30 min", L, EXT, "supported", "{2hx.pers30}", 12.14, 2,
          manifest="two-hour-external | persistence 30 / 60"),
        C("C-2X-11", "External two-hour persistence 20.44 at 60 min", L, EXT, "supported", "{2hx.pers60}", 20.44, 2),
        C("C-2X-12", "External two-hour ranges covered 87 %", L, EXT, "supported", "{2hx.cov120}*100", 87, 0,
          "Wider than needed (over-coverage)", manifest="two-hour-external | coverage 120"),
        C("C-CL-01", "Classifier warned 329 of 329 onsets (D-044)", "README §7; reports/two_hour_classifier.md", NEST, "negative-result",
          "{clf.detected_onsets}", 329, 0, "Warned at ~99 % of eligible time: degenerate", manifest="classifier-negative | onsets warned"),
        C("C-CL-02", "Classifier 0.18 false alarms per day", "README §7", NEST, "negative-result", "{clf.false_alarms_per_observed_day}", 0.18, 2,
          manifest="classifier-negative | false alarms per day"),
        C("C-CL-03", "Classifier warning active at 99 % of eligible time", "README §7", NEST, "negative-result",
          "{clf.share_of_eligible_time_warning}*100", 99, 0, manifest="classifier-negative | share of eligible time warning"),
        C("C-CL-04", "Classifier ROC AUC 0.64", "README §7", NEST, "negative-result", "{clf.origin_roc_auc}", 0.64, 2,
          manifest="classifier-negative | ROC AUC"),
        C("C-46-01", "Two-hour coverage before pooled calibration 76.5 %", L, NEST + " — re-analysis", "supported",
          "{wr.b0.cov120}*100", 76.5, 1, manifest="d046-ranges | coverage 120 before"),
        C("C-46-02", "Pooled calibration brought two-hour coverage to 79.8 % (D-046 B)", L, NEST + " — re-analysis", "supported",
          "{wr.b1.cov120}*100", 79.8, 1, "Not applied to the demo model", manifest="d046-ranges | coverage 120 after"),
        C("C-46-03", "60-min coverage 78.2 → 79.5 % (before)", L, NEST, "supported", "{wr.b0.cov60}*100", 78.2, 1,
          manifest="d046-ranges | coverage 60 before → after"),
        C("C-46-04", "60-min coverage 78.2 → 79.5 % (after)", L, NEST, "supported", "{wr.b1.cov60}*100", 79.5, 1),
        C("C-46-05", "Pooled ranges 7 mg/dL wider at 120 min", L, NEST, "supported", "{wr.b1.width120}-{wr.b0.width120}", 7, 0,
          manifest="d046-ranges | width increase"),
        C("C-46-06", "Per-person median coverage 81.2 % at 120 min", "README §7", NEST, "supported", "{wr.b1.pmed120}*100", 81.2, 1,
          manifest="d046-ranges | per-person median coverage 120"),
        C("C-46-07", "Earlier warning: 126 vs 123 of 329 warned ≥ 30 min (selected vs incumbent)", "README §7", NEST, "negative-result",
          "{wr.early.sel}", 126, 0, manifest="d046-earlier-warning | warned ≥ 30 min ahead, selected vs incumbent"),
        C("C-46-08", "Incumbent warned 123 of 329 ≥ 30 min ahead", "README §7", NEST, "supported", "{wr.early.inc}", 123, 0),
        C("C-46-09", "Earlier-warning difference +0.9 points", "README §7; docs/decisions.md D-046", NEST, "negative-result",
          "{wr.diff.d}*100", 0.9, 1, "CI [−5.5, +6.6] crosses zero", manifest="d046-earlier-warning | difference (95 % CI)"),
        C("C-46-10", "Earlier-warning CI lower −5.5 points", "README §7", NEST, "supported", "{wr.diff.lo}*100", -5.5, 1),
        C("C-46-11", "Earlier-warning CI upper +6.6 points", "README §7", NEST, "supported", "{wr.diff.hi}*100", 6.6, 1),
        C("C-46-12", "Median lead still 15 min (D-046 A)", "README §7", NEST, "supported", "{wr.lead.sel}", 15, 0,
          manifest="d046-earlier-warning | median lead"),
        C("C-46-13", "Profile prior −0.53 mg/dL at 120 min (D-046 C)", "README §7", NEST, "negative-result", "{wr.prior.all.d}", -0.53, 2,
          "CI [−1.57, +0.58] crosses zero", manifest="d046-profile-prior | 120 min"),
        C("C-46-14", "Profile prior first 24 h −0.93 mg/dL", "README §7", NEST, "negative-result", "{wr.prior.d1.d}", -0.93, 2,
          "CI [−2.74, +0.48]", manifest="d046-profile-prior | 120 min, first 24 h"),
        C("C-SC-01", "145 participants imported", "README §7", "Data inventory", "supported", "{inv.cgm}+{inv.sh}", 145, 0,
          manifest="scale | imported"),
        C("C-SC-02", "14 evaluated internally", "README §7", "Data inventory", "supported", "{inv.cgm_eval}", 14, 0,
          manifest="scale | evaluated internally"),
        C("C-SC-03", "98 evaluated externally", "README §7", "Data inventory", "supported", "{inv.sh_eval}", 98, 0,
          manifest="scale | evaluated externally"),
        C("C-50-01", "D-050 candidate pooled ranges 77.9 % at 30 min (test, retrospective)", "README §7; reports/serving_ranges_study.md",
          "Serving pipeline, final test (inspected) — candidate, not served", "supported", "{rg.pool.cov30}*100", 77.9, 1,
          "Research on inspected data; not the served interval", manifest="ranges-d050 | candidate coverage 30 / 60 (G on validation + calibration, test, retrospective)"),
        C("C-50-02", "D-050 candidate pooled ranges 78.7 % at 60 min", "README §7", "Candidate, not served", "supported", "{rg.pool.cov60}*100", 78.7, 1),
        C("C-50-03", "D-050 served coverage 74.2 % at 30 min", "README §7", SERV, "supported", "{rg.inc.cov30}*100", 74.2, 1,
          manifest="ranges-d050 | served coverage 30 / 60"),
        C("C-50-04", "D-050 served coverage 74.0 % at 60 min", "README §7", SERV, "supported", "{rg.inc.cov60}*100", 74.0, 1),
        C("C-50-05", "D-050: served ranges stay (trend-stratified shape not adopted)", "README §7; docs/decisions.md D-050", "Gate",
          "negative-result", None, None, None, "Pre-test coverage-error gain below the 1.0-point bar", sources="Gates: D-050; rg.chosen",
          manifest="ranges-d050 | decision"),
        C("C-51-01", "D-051 C0 60-min MAE 22.27", "README §7; reports/tune_study.md", NEST, "supported", "{tu.C0.mae60}", 22.27, 2,
          manifest="tune-d051 | C0 / C1 60-min MAE"),
        C("C-51-02", "D-051 C1 60-min MAE 21.90", "README §7", NEST, "supported", "{tu.C1.mae60}", 21.90, 2),
        C("C-51-03", "D-051 C1 − C0 −0.37 mg/dL (small internal gain vs equally tuned C0)", "README §7", NEST, "negative-result",
          "{tu.c1c0.mean_diff}", -0.37, 2, "Does not reach the declared −0.5 bar; not a release", manifest="tune-d051 | C1 − C0 (internal)"),
        C("C-51-04", "D-051 C1 − C0 CI [−0.63, −0.06] (lower)", "README §7", NEST, "supported", "{tu.c1c0.lo}", -0.63, 2),
        C("C-51-05", "D-051 C1 − C0 CI [−0.63, −0.06] (upper)", "README §7", NEST, "supported", "{tu.c1c0.hi}", -0.06, 2),
        C("C-51-06", "D-051 10 of 14 people better", "README §7", NEST, "supported", "{tu.c1c0.improved}", 10, 0, manifest="tune-d051 | people better"),
        C("C-51-07", "D-051 C1 − refitted serving recipe −0.65 [−1.41, +0.07] (not established)", "README §7", NEST, "negative-result",
          "{tu.c1r0.mean_diff}", -0.65, 2, "Different comparator from C0; CI crosses zero",
          manifest="tune-d051 | C1 − refitted serving recipe (not established)"),
        C("C-51-08", "D-051 not adopted", "README §7; docs/decisions.md D-051", "Gate", "negative-result", None, None, None,
          sources="Gates: D-051; tu.adopt.c1_better", manifest="tune-d051 | adoption"),
        C("C-52-01", "D-052 S 60-min personalised MAE 21.62", "reports/joint_study.md; docs/decisions.md D-052", NEST, "supported",
          "{jt.S.mae60}", 21.62, 2, manifest="joint-d052-forecast | S 60-min personalised MAE"),
        C("C-52-02", "D-052 R0 60-min personalised MAE 22.50 (refitted serving recipe)", "reports/joint_study.md", NEST, "supported",
          "{jt.R0.mae60}", 22.50, 2, manifest="joint-d052-forecast | R0 60-min personalised MAE (refitted serving recipe)"),
        C("C-52-03", "D-052 S − R0 −0.88 [−1.75, −0.13] at 60 min", "reports/joint_study.md", NEST, "supported", "{jt.gate.mean_diff_60}", -0.88, 2,
          manifest="joint-d052-forecast | S − R0 at 60 min (participant bootstrap)"),
        C("C-52-04", "D-052 primary gate not met (high-target +1.70 vs +1.0)", "docs/decisions.md D-052", "Gate", "negative-result",
          "{jt.gate.high_worsening_60}", 1.70, 2, manifest="joint-d052-forecast | primary gate"),
        C("C-52-05", "D-052 release guardrails not met", "docs/decisions.md D-052", "Gate", "negative-result", None, None, None,
          sources="Gates: D-052", manifest="joint-d052-operational | release guardrails vs R0"),
        C("C-52-06", "D-052 candidate 80 % range coverage 78.4 % at 60 min", "docs/decisions.md D-052", NEST, "supported", "{jt.cov60}*100", 78.4, 1,
          "Mixed-fit calibration erratum disclosed in reports/joint_study.md", manifest="joint-d052-operational | candidate 80 % range coverage, 60 min"),
        C("C-53-01", "D-053 S 60-min personalised MAE 22.00", RT + "; README §7", NEST, "supported", "{cs.S.mae60}", 22.00, 2,
          manifest="constrained-d053 | S 60-min personalised MAE"),
        C("C-53-02", "D-053 S 30-min MAE 13.89", RT, NEST, "supported", "{cs.S.mae30}", 13.89, 2),
        C("C-53-03", "D-053 persistence 15.34 / 24.47 (30 min)", RT, NEST, "supported", "{cs.persistence.mae30}", 15.34, 2),
        C("C-53-04", "D-053 persistence 15.34 / 24.47 (60 min)", RT, NEST, "supported", "{cs.persistence.mae60}", 24.47, 2),
        C("C-53-05", "D-053 R0 refitted in the same folds 23.20", RT, NEST, "supported", "{cs.R0.mae60}", 23.20, 2,
          manifest="constrained-d053 | R0 60-min personalised MAE (refitted in the same folds)"),
        C("C-53-06", "D-053 S − R0 −1.19 at 60 min", RT, NEST, "supported", "{cs.p60.mean_diff}", -1.19, 2,
          manifest="constrained-d053 | S − R0 at 60 min (participant bootstrap)"),
        C("C-53-07", "D-053 S − R0 CI lower −2.09", RT, NEST, "supported", "{cs.p60.lo}", -2.09, 2),
        C("C-53-08", "D-053 S − R0 CI upper −0.39", RT, NEST, "supported", "{cs.p60.hi}", -0.39, 2),
        C("C-53-09", "D-053 lower for 11 of 14 people at 60 min", "README §7", NEST, "supported", "{cs.p60.improved}", 11, 0,
          manifest="constrained-d053 | people better at 60 min"),
        C("C-53-10", "D-053 30-min change −0.50", "README §7", NEST, "supported", "{cs.pg.worsening_30}", -0.50, 2,
          manifest="constrained-d053 | 30-min change"),
        C("C-53-11", "D-053 error above 180 rose 1.07 against a +1.0 limit", RT, NEST, "negative-result", "{cs.pg.high_worsening_60}", 1.07, 2,
          "Unrounded +1.0658 > +1.0: gate fails at any display rounding", manifest="constrained-d053 | high-target change at 60 min"),
        C("C-53-12", "D-053 primary gate not met; not adopted", RT + "; docs/decisions.md D-053", "Gate", "negative-result", None, None, None,
          sources="Gates: D-053; cs.pg.passed", manifest="constrained-d053 | primary gate"),
        C("C-53-13", "D-053 false alarms per observed day S 1.15", "README §7", NEST, "supported", "{cs.w.S.fa_day}", 1.15, 2,
          manifest="constrained-d053 | false alarms per observed day, S"),
        C("C-53-14", "D-053 false alarms per observed day R0-P 1.20", "README §7", NEST, "supported", "{cs.w.R0P.fa_day}", 1.20, 2,
          manifest="constrained-d053 | false alarms per observed day, R0-P"),
        C("C-53-15", "D-053 onsets warned ≥ 30 min ahead, S 21.9 %", "README §7", NEST, "supported", "{cs.w.S.share30}*100", 21.9, 1,
          manifest="constrained-d053 | onsets warned ≥ 30 min ahead, S"),
        C("C-53-16", "D-053 onsets warned ≥ 30 min ahead, R0-P 16.4 %", "README §7", NEST, "supported", "{cs.w.R0P.share30}*100", 16.4, 1,
          manifest="constrained-d053 | onsets warned ≥ 30 min ahead, R0-P"),
        C("C-53-17", "D-053 operational gate passed", "README §7", "Gate", "supported", None, None, None, sources="Gates: D-053 operational",
          manifest="constrained-d053 | operational gate"),
        C("C-53-18", "D-053 range coverage 79.4 % at 30 min", "README §7", NEST, "supported", "{cs.cov30}*100", 79.4, 1,
          manifest="constrained-d053 | range coverage 30 min"),
        C("C-53-19", "D-053 range coverage 77.3 % at 60 min", "README §7", NEST, "supported", "{cs.cov60}*100", 77.3, 1,
          manifest="constrained-d053 | range coverage 60 min"),
        C("C-54-01", "D-054 foundation 60-min MAE 22.92", RT, NEST + " — all origins", "supported", "{fd.FM60}", 22.92, 2,
          "Zero-shot, CGM only, not fine-tuned, not served", manifest="foundation-d054 | foundation 60-min MAE"),
        C("C-54-02", "D-054 foundation 30-min MAE 14.55", RT, NEST, "supported", "{fd.FM30}", 14.55, 2),
        C("C-54-03", "D-054 persistence 24.47 on identical origins", RT, NEST, "supported", "{fd.P60}", 24.47, 2,
          manifest="foundation-d054 | persistence 60-min MAE (identical origins)"),
        C("C-54-04", "D-054 foundation − persistence −1.54 [−1.81, −1.30]", "README §7", NEST, "supported", "{fd.fp.mean}", -1.54, 2,
          manifest="foundation-d054 | foundation − persistence"),
        C("C-54-05", "D-054 all 14 people better than persistence", "README §7", NEST, "supported", "{fd.fp.people_better}", 14, 0,
          manifest="foundation-d054 | people better than persistence"),
        C("C-54-06", "D-054 refitted serving recipe 22.50 at 60 min", RT, NEST, "supported", "{fd.R060}", 22.50, 2,
          manifest="foundation-d054 | refitted serving recipe 60-min MAE (saved D-052 R0)"),
        C("C-54-07", "D-054 foundation − refitted recipe +0.42 [−0.70, +1.47]: not better", "README §7", NEST, "negative-result",
          "{fd.fr.mean}", 0.42, 2, manifest="foundation-d054 | foundation − refitted recipe"),
        C("C-54-08", "D-054 F1 residual − same-code refit −0.46 [−0.72, −0.21] (descriptive)", "README §7", NEST, "supported",
          "{fd.f1.mean}", -0.46, 2, "Descriptive secondary evidence; below the −0.5 bar used elsewhere; not promoted",
          manifest="foundation-d054 | F1 residual − same-code refit (descriptive)"),
        C("C-PR-01", "App replay reproduced all 15,440 saved test forecasts", RT + "; README Safety guardrails", "Software test (saved parity run)",
          "tested", "{parity.compared}", 15440, 0, "15,440 = forecasts (7,720 issued origins × 2 horizons), not origins or people",
          manifest="parity | forecasts compared (saved parity run, outcome pass for the serving model)"),
        C("C-PR-02", "Parity tolerance 0.05 mg/dL", RT, "Software test", "tested", "{parity.tolerance_mg_dl}", 0.05, 2,
          manifest="parity | tolerance (mg/dL)"),
        C("C-SY-01", "Synthetic demonstration model is fitted on fictional data, never evidence", "README (Public demonstration)",
          "Synthetic mode", "supported", None, None, None, sources="reg.synthetic_fitted; registry",
          manifest="synthetic | synthetic-demo.lightgbm.dyn-v2.v1 fitted on fictional data (registry synthetic_fitted)"),
    ]
    return out


def readme_claims() -> list[dict]:
    R = "README.md §7"
    return [
        C("C-FC-05", "Serving vs persistence −1.88 mg/dL at 30 min", R, SERV, "supported", "{fus.cmp.sp.d30}", -1.88, 2,
          "Candidate − control; participant-cluster bootstrap"),
        C("C-FC-06", "Serving vs persistence 30-min CI [−2.53, −1.32]", R, SERV, "supported", "{fus.cmp.sp.hi30}", -1.32, 2, sources="fus.cmp.sp.lo30"),
        C("C-FC-07", "Serving vs persistence −2.96 [−4.34, −1.77] at 60 min", R, SERV, "supported", "{fus.cmp.sp.d60}", -2.96, 2),
        C("C-FC-08", "Serving better than persistence for 14/14 participants at 30 min", R, SERV, "supported", "{fus.better30}", 14, 0),
        C("C-FC-09", "Serving better than persistence for 13/14 participants at 60 min", R, SERV, "supported", "{fus.better60}", 13, 0),
        C("C-FC-10", "Patient adaptation −0.18 [−0.35, −0.01] at 30 min", R, SERV, "supported", "{fus.cmp.adapt.d30}", -0.18, 2,
          "Small"),
        C("C-FC-11", "Patient adaptation −0.33 [−0.68, +0.02] at 60 min: inconclusive", R, SERV, "negative-result", "{fus.cmp.adapt.d60}", -0.33, 2,
          "CI crosses zero"),
        C("C-FC-12", "LightGBM minus CatBoost +0.46 [+0.05, +0.92] at 60 min (population)", R, SERV, "supported", "{fus.cmp.lgbcb.d60}", 0.46, 2,
          "CatBoost better at 60 min; tie-break chose LightGBM"),
        C("C-FC-13", "Profile vs no profile within the same 14 people −0.25 at 30 min", R, SERV, "supported", "{fus.cmp.prof.d30}", -0.25, 2,
          "May be identity, not physiology"),
        C("C-FC-14", "Profile vs no profile within the same 14 people −0.61 at 60 min", R, SERV, "supported", "{fus.cmp.prof.d60}", -0.61, 2),
        C("C-FC-15", "CatBoost + profile + adaptation 12.93 / 19.80 (tied on validation, not selected) — 30 min", R, SERV, "supported",
          "{fus.cbpa.mae30}", 12.93, 2),
        C("C-FC-16", "CatBoost + profile + adaptation 19.80 at 60 min", R, SERV, "supported", "{fus.cbpa.mae60}", 19.80, 2),
        C("C-FC-17", "Previous demo model (v2) 13.15 / 20.81 — 30 min", R, SERV, "supported", "{fus.v2a.mae30}", 13.15, 2),
        C("C-FC-18", "Previous demo model (v2) 20.81 at 60 min", R, SERV, "supported", "{fus.v2a.mae60}", 20.81, 2),
        C("C-FC-19", "7,730 forecast origins in the final test segment", R, SERV, "supported", "{fus.seg.test.origins}", 7730, 0,
          "Origins ≠ issued origins (7,720) ≠ matched targets"),
        C("C-FC-20", "≈ 644 h of final-test recordings", R, SERV, "supported", "{fus.seg.test.covered_hours}", 644, 0),
        C("C-WN-07", "Precision 0.60", R, SERV, "supported", "{fus.ev.precision}", 0.60, 2),
        C("C-WN-08", "2.28 false alarms per eligible day", R, SERV, "supported", "{fus.ev.false_alarm_episodes}/{fus.ev.eligible_patient_days}",
          2.28, 2, "Eligible patient-days differ from observed patient-days"),
        C("C-WN-09", "48.4 % of all onsets warned at least 30 min ahead", RT + "; reports/warning_utility.md", SERV, "supported",
          "{wu.ge30}/{wu.onsets}*100", 48.4, 1, "Denominator = all 64 eligible onsets (missed onsets count against it)"),
        C("C-WN-10", "Warned ≥ 15 min ahead 67.2 % of all onsets", "reports/warning_utility.md", SERV, "supported",
          "{wu.ge15}/{wu.onsets}*100", 67.2, 1),
        C("C-WN-11", "Warned ≥ 60 min ahead 20.3 % of all onsets", "reports/warning_utility.md", SERV, "supported",
          "{wu.ge60}/{wu.onsets}*100", 20.3, 1),
        C("C-WN-12", "Qualified lead IQR 10–55 min (detected onsets)", "reports/warning_utility.md", SERV, "supported",
          "{wu.lead_q75}", 55, 0, sources="wu.lead_q25"),
        C("C-WN-13", "40 of 101 warning episodes were false alarms", "README (GIF caption)", SERV, "supported", "{fus.ev.warning_episodes}", 101, 0),
        C("C-WN-14", "Persistence-based warning: 63/64 detected with 70 false alarms (2.61 per observed day)", R, SERV, "supported",
          "{fus.pev.false_alarms_per_observed_patient_day}", 2.61, 2),
        C("C-RG-05", "Ranges covered only 60–82 % (30 min) per participant (min)", R, SERV, "supported", "{fus.iv.per_patient_min30}*100", 60, 0),
        C("C-RG-06", "Ranges per participant 30-min max 82 %", R, SERV, "supported", "{fus.iv.per_patient_max30}*100", 82, 0),
        C("C-RG-07", "Ranges per participant 60-min 51–88 % (min)", R, SERV, "supported", "{fus.iv.per_patient_min60}*100", 51, 0),
        C("C-RG-08", "Ranges per participant 60-min max 88 %", R, SERV, "supported", "{fus.iv.per_patient_max60}*100", 88, 0),
        C("C-RG-09", "While glucose rose (> 1 mg/dL/min) served ranges covered 51.6 % at 60 min on test", "docs/decisions.md D-050",
          SERV, "supported", "{rg.inc.rising60}*100", 51.6, 1, "Final-test trend context (read once in D-050)"),
        C("C-TR-01", "Validation: 60-min error 35.3 mg/dL while glucose rose > 2 mg/dL/min", "README §5", "Serving model, validation segment",
          "supported", "{tr.rising_fast.mae60}", 35.3, 1, "Validation segment, not the test"),
        C("C-TR-02", "Validation: 60-min error 19.5 when steady", "README §5", "Serving model, validation segment", "supported",
          "{tr.steady.mae60}", 19.5, 1),
        C("C-TR-03", "Peaks above 180 under-forecast by 16.9 mg/dL at 60 min (validation)", "README §5", "Serving model, validation segment",
          "supported", "-{tr.bias_high60}", 16.9, 1),
        C("C-EX-06", "External 30-min difference −0.99 [−1.21, −0.77], better for 80 of 98", R, EXT, "supported", "{ext.cmp.mean_diff30}", -0.99, 2),
        C("C-EX-07", "External better for 80 of 98 at 30 min", R, EXT, "supported", "{ext.cmp.improved30}", 80, 0),
        C("C-EX-08", "External 60-min difference −0.34 [−0.76, +0.10]: inconclusive", R, EXT, "negative-result", "{ext.cmp.mean_diff60}", -0.34, 2),
        C("C-EX-09", "External ranges covered 87.1 % / 86.1 % (30 min)", R, EXT, "supported", "{ext.cov30}*100", 87.1, 1),
        C("C-EX-10", "External ranges 86.1 % at 60 min", R, EXT, "supported", "{ext.cov60}*100", 86.1, 1),
        C("C-NS-06", "v2 recipe refit per fold 14.52 / 24.25 (30 min)", R, NEST, "supported", "{nest.v2.mae30}", 14.52, 2),
        C("C-NS-07", "v2 recipe refit per fold 24.25 at 60 min", R, NEST, "supported", "{nest.v2.mae60}", 24.25, 2),
        C("C-2H-12", "Two-hour selected minus persistence −5.25 [−6.44, −4.05]", "docs/evidence_checklist_2026-10-20.md", NEST, "supported",
          "{2h.cmp.d}", -5.25, 2),
        C("C-GD-01", "No forecast is issued with less than 60 min of history (abstains)", "README Safety guardrails", "Serving software",
          "tested", sources="tests.test_replay_api::test_insufficient_history_abstains_with_nulls"),
        C("C-GD-02", "After 15 min without a reading the last forecast is shown marked Stale", "README Safety guardrails", "Serving software",
          "tested", "{cfg.stale}", 15, 0, sources="e2e: Participant 12: stale glucose before the gap closes, then abstention"),
        C("C-GD-03", "Correction needs ≥ 6 matured outcomes, capped at ±25 mg/dL, resets after a 3-hour gap", "README Safety guardrails",
          "Serving software", "tested", "{cfg.reset_gap}/60", 3, 0,
          sources="tests.test_replay_api::test_personalization_reports_its_bound_and_reset_gap"),
        C("C-GD-04", "Every warning is labelled Research warning and stays Outcome pending until it can be judged",
          "README Safety guardrails", "Serving software", "tested",
          sources="e2e: Participant 01: a warning is pending until its outcome is knowable"),
        C("C-GD-05", "Opened tests write a lock file; re-running needs a written reason", "README Safety guardrails", "Research software",
          "tested", sources="tests.test_research_pipeline::test_editing_a_frozen_protocol_is_refused"),
        C("C-GD-06", "Event = new episode above 180 mg/dL sustained ≥ 15 minutes", "README §4", "Event definition", "implemented",
          "{cfg.threshold}", 180, 0, sources="cfg.confirm; tests.test_events::test_exact_cutoff_is_not_high"),
        C("C-AU-01", "CI green on both branches", "context.md (handoff)", "CI", "author-reported", None, None, None,
          "No saved CI log in the repository", "Link the GitHub Actions run after the repository is public"),
        C("C-AU-02", "Fresh-clone checks on 23 Oct: pytest pass, vitest 141/141, e2e 16 passed / 13 skipped (synthetic)",
          "Team handoff (23 Oct)", "Fresh clone, synthetic data", "author-reported", None, None, None,
          "Logs not saved; not independently verified", "Save a ledger for the fresh-clone run, or cite only the f073a10 ledger"),
        C("C-VR-01", "pytest 221 passed, 2 skipped at f073a10", "context.md; reports/verification/final_2026-10-22_145326Z/LEDGER.md",
          "Local verification at f073a10 (not later revisions)", "tested", "{ver.pytest.passed}", 221, 0, sources="ver.pytest.skipped"),
        C("C-VR-02", "vitest 141 passed at f073a10", "context.md; LEDGER.md", "Local verification at f073a10", "tested", "{ver.vitest.passed}", 141, 0),
        C("C-VR-03", "e2e 29 normal/recorded at f073a10", "context.md; LEDGER.md", "Local verification at f073a10", "tested", "{ver.e2e.passed}", 29, 0),
        C("C-VR-04", "e2e 4 public at f073a10", "LEDGER.md", "Local verification at f073a10", "tested", "{ver.e2e_public.passed}", 4, 0),
        C("C-VR-05", "e2e 3 summary/export at f073a10", "LEDGER.md", "Local verification at f073a10", "tested", "{ver.e2e_summary.passed}", 3, 0),
        C("C-VR-06", "Parity re-run in the final ledger: 15,440 compared, 0 over tolerance", "LEDGER.md", "Local verification at f073a10",
          "tested", "{parity_final.compared}", 15440, 0, sources="parity_final.over_tolerance"),
    ]


# Frozen slides 1–5 (visible text, 24 Oct freeze). Status never upgraded without a primary source.
def frozen_slide_claims() -> list[dict]:
    return [
        C("C-S1-01", "GlucoTwin — Team Kryptonite, VESIT; Avinash Gehi (team leader) · Hardik Hinduja", "Slide 1", "Team facts",
          "supported", sources="submission (configs/submission.yaml team.*)"),
        C("C-S1-02", "Digital Twin Challenge 2026 · Happiest Health", "Slide 1; README", "Challenge identity", "author-reported",
          limitation="Challenge page not archived in the repository", action="Optional: cite the Unstop challenge URL"),
        C("C-S1-03", "Student research prototype — not a medical device", "Slide 1; README", "Scope statement", "author-reported",
          limitation="A declaration of intended use; no regulatory submission exists or is claimed"),
        C("C-S2-01", "HbA1c summarises ~3 months of glucose; it cannot show short high-glucose excursions (NIDDK)", "Slide 2; README §4",
          "External fact", "pending", limitation="No NIDDK URL is cited anywhere in the repository",
          action="Pending primary-source check: locate the NIDDK page and record its URL and access date"),
        C("C-S2-02", "WHO identifies continuity of care and diabetes monitoring as ongoing challenges in India", "Slide 2",
          "External fact", "pending", limitation="No WHO document is cited anywhere in the repository",
          action="Pending primary-source check: locate the WHO source and record its URL and access date"),
        C("C-S2-03", "Continuous glucose monitors produce ~288 readings a day", "Slide 2; README §4–5", "Arithmetic",
          "supported-arithmetic", "{const.min_per_day}/{const.cadence}", 288, 0, "Assumes a 5-minute sensor cadence (Dexcom G6 Pro)"),
        C("C-S2-04", "Question for the clinician (framing)", "Slide 2", "Framing", "proposed/illustrative"),
        C("C-S2-05", "Target: forecast glucose 30 and 60 minutes ahead and warn of a NEW sustained episode > 180 mg/dL (≥ 15 min), Type 2",
          "Slide 2; README §4", "Task definition", "implemented", "{cfg.confirm}", 15, 0, sources="cfg.threshold; fus.participants"),
        C("C-S3-01", "Hypothetical Mumbai diabetes educator illustration", "Slide 3; README §5", "Illustration", "proposed/illustrative",
          limitation="No clinic, partner, pilot or Indian participants"),
        C("C-S3-02", "Proposed users and workflow", "Slide 3", "Proposal", "proposed/illustrative"),
        C("C-S3-03", "Never: no insulin advice, triage, live device claims or LLM predictions", "Slide 3; README", "Scope boundary",
          "implemented", limitation="Absence of features verified by code inspection, not by a dedicated test"),
        C("C-S3-04", "Potential fit with diagnostics + outpatient services is a proposal — no partnership or pilot is claimed", "Slide 3",
          "Proposal", "proposed/illustrative"),
        C("C-S4-01", "Versioned patient state", "Slide 4", "Serving software", "tested",
          sources="tests.test_replay_api::test_review_snapshot_is_one_session_revision_and_clock"),
        C("C-S4-02", "Replay clock releases only what was available", "Slide 4", "Serving software", "tested",
          sources="tests.test_replay_api::test_event_polling_has_no_duplicates_or_gaps_and_respects_availability; tests.test_leakage::test_late_logged_meal_is_hidden_until_available"),
        C("C-S4-03", "Population forecaster + patient correction learned from matured errors", "Slide 4", "Serving software", "tested",
          sources="tests.test_leakage::test_sixty_minute_residual_matures_only_when_target_is_available; tests.test_leakage::test_corrections_never_stack_on_corrected_predictions"),
        C("C-S4-04", "A statistical research twin, not a physiological simulator", "Slide 4", "Definition", "implemented",
          limitation="Describes the method class (no physiological model in code)"),
        C("C-S5-01", "CGMacros 1.0.0 (PhysioNet, CC BY-NC-SA 4.0)", "Slide 5; docs/data_card.md", "Dataset licence", "supported",
          sources="src.cgm.license; src.cgm.doi"),
        C("C-S5-02", "CGMacros: 45 adults", "Slide 5", "Dataset", "supported", "{inv.cgm}", 45, 0),
        C("C-S5-03", "~10 days, Dexcom G6 Pro + Libre Pro (blinded), Fitbit, meals, screening labs", "Slide 5; docs/data_card.md",
          "Dataset description", "supported", sources="data_card",
          limitation="From the publisher description recorded in the data card; not re-checked against PhysioNet in this build"),
        C("C-S5-04", "Primary cohort: 14 participants in the Type 2 group (HbA1c > 6.4 %)", "Slide 5; README", "Cohort", "supported",
          "{fus.participants}", 14, 0, sources="decisions (D-006)"),
        C("C-S5-05", "31 others replayable but excluded", "Slide 5; docs/data_card.md", "Cohort", "supported", "{inv.cgm}-{fus.participants}",
          31, 0, sources="fus.excluded"),
        C("C-S5-06", "Published 1-minute CGM columns are linearly interpolated; native 5-minute Dexcom readings recovered", "Slide 5",
          "Data handling", "tested", sources="data_audit; tests.test_leakage::test_one_minute_interpolated_rows_do_not_inflate_observations"),
        C("C-S5-07", "Undocumented columns reported; dates shifted; timezone unknown → synthetic clock, no time-of-day features",
          "Slide 5", "Data handling", "supported", sources="data_audit (Undocumented columns; Time)",
          limitation="Data card says the synthetic timeline preserves source time of day; 'no time-of-day features' is a model-input statement checked against the feature list, not a test"),
        C("C-S5-08", "Synthetic fictional participants (labelled synthetic)", "Slide 5; README", "Synthetic mode", "tested",
          sources="tests.test_public_demo::test_synthetic_only; reg.synthetic_fitted"),
    ]


# Slides 6–12 and narration. Two ways to add:
#   1. A new claim:   DECK_CLAIMS.append(C("C-S9-01", "exact slide text", "Slide 9; narration 05:10", "<scope>",
#                                         "supported", "{fus.serv.mae30}", 12.96, 2, "limitation", "action"))
#      ``expr`` uses value ids from the Sources sheet ({id}); text-only claims use sources="..." instead of expr.
#   2. An existing claim that also appears on a slide or in the narration: EXTRA_LOCATIONS.append(("C-FC-01", "Slide 9"))
DECK_CLAIMS: list[dict] = []
EXTRA_LOCATIONS: list[tuple[str, str]] = []

D6, D7, D8, D9, D10, D11, D12, D13, D14, D15 = (f"Deck s{n}" for n in range(6, 16))
DECK_CLAIMS += [
    C("C-D6-01", "A forecast at time t uses only records observed and available by t", D6, "Serving software", "tested",
      sources="tests.test_leakage::test_future_values_cannot_change_earlier_features; tests.test_data_expansion::test_a_reading_released_late_is_invisible_before_its_availability"),
    C("C-D6-02", "Meals left out of the main model (logging time unknown); tested separately", D6, SERV, "supported",
      "{fus.sens.meal60}", None, None, "Optimistic meal-logging sensitivity, not eligible for serving", sources="fus.sens.meal60; data_card"),
    C("C-D6-03", "Protocol frozen before fitting; choices made on validation", D6, "Protocol", "supported",
      sources="Studies sheet (protocol first commit vs run); fus.selection_rule",
      limitation="The fusion protocol's first commit is 25 s after the run's created_at and its constraint was declared after the test was inspected (disclosed)"),
    C("C-D6-04", "The fusion re-run (gt-primary-2026-10-13-fusion) was declared after gt-primary-2026-10-11-v2's test had been inspected",
      D6 + "; " + D9, SERV, "supported", sources="decisions (D-038/D-039); fusion /correction_note"),
    C("C-D7-01", "Synthetic patients: 5 fictional profiles (CC0)", D7, "Synthetic mode", "implemented",
      sources="ml/glucotwin/data/synthetic.py (scenarios SYN-A…SYN-E, LICENSE CC0-1.0); docs/data_card.md"),
    C("C-D7-02", "External cohort ShanghaiT2DM, 98 people, 15-min CGM, research only", D7, EXT, "supported", "{inv.sh_eval}", 98, 0),
    C("C-D7-03", "Replay releases a reading once available_at ≤ t; seek / reset → new revision", D7, "Serving software", "tested",
      sources="tests.test_replay_api::test_event_polling_has_no_duplicates_or_gaps_and_respects_availability; tests.test_replay_api::test_seek_backward_removes_future_state_and_invalidates_cursors"),
    C("C-D7-04", "Causal features, ≤ 180 min history", D7, "Serving software", "implemented", "{cfg.lookback}", 180, 0),
    C("C-D7-05", "Fusion LightGBM: 8 profile + 38 dynamic features", D7 + "; " + D8, "Serving model", "author-reported",
      limitation="Counts stated in docs/model_card.md and README; not re-derived from the model artifact by this builder",
      action="Optional: read the feature list from the registered model bundle"),
    C("C-D7-06", "Split 60/10/10/20 with purged boundaries", D7, "Protocol", "tested",
      sources="tests.test_leakage::test_split_purging_keeps_label_support_inside_segments; configs/experiment.yaml"),
    C("C-D7-07", "CPU only, local, no cloud, no LLM; not a medical device", D7, "Scope boundary", "implemented",
      limitation="Absence verified by code inspection only"),
    C("C-D8-01", "Baselines persistence, linear trend, ridge; learners LightGBM and CatBoost, same search budget", D8, SERV, "supported",
      sources="fus.selection_rule; selection candidates (Sources sel.*)"),
    C("C-D8-02", "Patient correction: a bounded mean of this person's observed errors; model weights never refit", D8,
      "Serving software", "tested", sources="tests.test_leakage::test_corrections_never_stack_on_corrected_predictions; tests.test_replay_api::test_personalization_reports_its_bound_and_reset_gap"),
    C("C-D8-03", "80 % prediction ranges (split-conformal), calibrated for the exact serving pipeline", D8, SERV, "supported",
      "{const.target_cov}*100", 80, 0, sources="fus.iv.q30; fus.iv.q60"),
    C("C-D8-04", "Warning score: highest forecast in the next 60 min minus 180 mg/dL — not a probability", D8, "Serving software",
      "implemented", "{cfg.threshold}", 180, 0),
    C("C-D9-10", "Rule: lowest mean per-participant validation MAE over both horizons; within 0.5 mg/dL → simpler model", D9, SERV,
      "supported", sources="fus.selection_rule"),
    C("C-D9-11", "Only dyn-prof-v2 eligible (fusion constraint added 2026-10-13, after the test segment had been inspected)", D9, SERV,
      "supported", sources="fus.selection_rule; decisions (D-038)"),
    C("C-D9-12", "Selected: lightgbm:dyn-prof-v2 with patient correction", D9, SERV, "supported", sources="sel.selected; reg.serving"),
    C("C-D9-13", "Without the constraint the rule picks lightgbm:dyn-v2 (gt-primary-2026-10-11-v2)", D9, SERV, "supported",
      sources="sel.v2_selected"),
    C("C-D12-01", "amazon/chronos-bolt-tiny, 8.7 M parameters, apache-2.0, zero-shot on CGM only", D12, "Research baseline (D-054)",
      "supported", "{fd.params}/1000000", 8.7, 1, sources="fd.checkpoint; fd.license"),
    C("C-D12-02", "Foundation vs persistence −0.79 [−1.01, −0.59] mg/dL at 30 min", D12, NEST, "supported", "{fd.fp30.mean}", -0.79, 2),
    C("C-D12-03", "30-min CI lower −1.01", D12, NEST, "supported", "{fd.fp30.lo}", -1.01, 2),
    C("C-D12-04", "30-min CI upper −0.59", D12, NEST, "supported", "{fd.fp30.hi}", -0.59, 2),
    C("C-D12-05", "Foundation vs persistence 60-min CI lower −1.81", D12, NEST, "supported", "{fd.fp.lo}", -1.81, 2),
    C("C-D12-06", "Foundation vs persistence 60-min CI upper −1.30", D12, NEST, "supported", "{fd.fp.hi}", -1.30, 2),
    C("C-D12-07", "Foundation vs refitted recipe CI [−0.70, +1.47] (lower)", D12, NEST, "supported", "{fd.fr.lo}", -0.70, 2),
    C("C-D12-08", "Foundation vs refitted recipe CI upper +1.47", D12, NEST, "supported", "{fd.fr.hi}", 1.47, 2),
    C("C-D12-09", "Foundation worse above 180 mg/dL than the refitted recipe", D12, NEST, "negative-result",
      "{fd.high.FM}-{fd.high.R0}", None, None, "Positive difference = foundation worse", sources="fd.high.FM; fd.high.R0"),
    C("C-D12-10", "Its pretraining data may include public CGM recordings", D12, "Disclosure", "author-reported",
      limitation="A caution, not a measured fact"),
    C("C-D13-01", "Built: replay API, immutable forecasts, warnings with outcomes, clinician workspace", D13, "Serving software", "tested",
      sources="tests.test_replay_api::test_forward_replay_equals_reset_and_seek; e2e: Participant 01: a warning is pending until its outcome is knowable"),
    C("C-D13-02", "Reproducible: synthetic mode needs no download; results regenerate from versioned configs", D13, "Software", "tested",
      sources="tests.test_replay_api::test_synthetic_demo_runs_without_dataset"),
    C("C-D14-01", "14 participants from one US site; no Indian participants", D14, "Limitation", "supported", "{fus.participants}", 14, 0,
      sources="data_card"),
    C("C-D14-02", "Blinded CGMs replayed with assumed zero latency; medication and meal timing unknown", D14, "Limitation", "supported",
      sources="data_card; docs/model_card.md"),
    C("C-D14-03", "Ranges under-cover (74.2 % / 74.0 % vs 80 % target)", D14 + "; " + D11, SERV, "supported",
      "({fus.iv.coverage30}-{const.target_cov})*100", None, None, "Negative = under-coverage", sources="fus.iv.coverage60"),
    C("C-D14-04", "The warning threshold ignores review workload", D14, "Limitation", "author-reported",
      limitation="Threshold chosen by event F1 on validation (implemented); the workload statement is a design observation"),
    C("C-D14-05", "Retrospective replay only — no clinical, workflow or safety study", D14, "Limitation", "supported",
      sources="README; docs/model_card.md"),
    C("C-D14-06", "Next: a prospective, consented Indian cohort and a full-fusion evaluation on new people", D14, "Proposal",
      "proposed/illustrative"),
    C("C-D15-01", "Code: MIT license (code only); quick start make setup · make synthetic · make dev", D15, "Repository", "supported",
      sources="submission (license.software MIT); LICENSE; Makefile"),
    C("C-D15-02", "Public GitHub link and demo video added after publication/upload", D15, "Submission status", "pending",
      action="Fill links.repository and links.video in configs/submission.yaml after the team publishes and uploads"),
    C("C-D15-03", "ShanghaiT2DM (figshare, doi:10.6084/m9.figshare.21600933.v5, CC BY 4.0)", D15, "Dataset licence", "supported",
      sources="src.sh.doi; src.sh.license"),
    C("C-NR-01", "Per-person 60-min range coverage 51.4 %–87.6 % (min)", "Narration metric card", SERV, "supported",
      "{fus.iv.per_patient_min60}*100", 51.4, 1),
    C("C-NR-02", "Per-person 60-min range coverage max 87.6 %", "Narration metric card", SERV, "supported",
      "{fus.iv.per_patient_max60}*100", 87.6, 1),
    C("C-NR-03", "Range held 46 % while rising > 2 mg/dL/min at 60 min (validation)", "Narration scene 15; metric card",
      "Serving model, validation segment", "supported", "{tr.rising_fast.cov60}*100", 46, 0),
    C("C-NR-04", "Joint selection lowered 60-min error but made peaks above 180 worse and raised false alarms: failed its gates",
      "Narration scene 11", NEST, "negative-result", "{jt.rel.fa_per_day_delta}", None, None, sources="Gates: D-052; jt.gate.high_worsening_60"),
    C("C-NR-05", "D-053 passed the warning and range checks", "Narration scene 11, 16", "Gate", "supported", sources="Gates: D-053 operational; cs.rg.passed"),
    C("C-NR-06", "Blending, retraining, tuning and joint selection each failed their declared bars", "Narration scene 16", "Gate",
      "negative-result", sources="Gates: D-048, D-049, D-051, D-052"),
    C("C-NR-07", "Shanghai: a separate CGM-only model beat persistence at 30 min and at two hours", "Narration scene 16", EXT, "supported",
      "{2hx.cmp.hi}", None, None, "120-min CI upper < 0; 30-min D-032 CI upper < 0", sources="ext.cmp.hi30"),
]
SLIDE9_STATED = {  # as printed on slide 9: val. MAE 30, val. MAE 60, score
    "lightgbm:dyn-prof-v2": (13.47, 21.26, 17.36), "catboost:dyn-prof-v2": (13.56, 21.31, 17.44),
    "lightgbm:dyn-v2": (13.76, 21.73, 17.75), "catboost:dyn-v2": (13.88, 21.66, 17.77), "ridge:dyn-v2": (14.13, 22.44, 18.28),
    "ridge:dyn-prof-v2": (14.21, 22.61, 18.41), "linear_trend:cgm-only-v1": (15.66, 26.33, 20.99),
    "persistence:cgm-only-v1": (16.51, 26.32, 21.42)}
for _slide9_i, _lab in enumerate(SLIDE9, 1):
    for _j, (_k, _d) in enumerate((("selection_mae_30", "val. MAE 30"), ("selection_mae_60", "val. MAE 60"),
                                   ("selection_score", "score"))):
        DECK_CLAIMS.append(C(f"C-D9-{_slide9_i:02d}{_k[-2:] if _k != 'selection_score' else 'sc'}", f"Slide 9 table: {_lab} {_d}",
                             D9, "Serving experiment, validation segment (selection)", "supported", f"{{sel.{_lab.replace(':', '_')}.{_k}}}",
                             SLIDE9_STATED[_lab][_j], 2))
EXTRA_LOCATIONS += [
    ("C-WN-04", D13), ("C-WN-01", D13), ("C-WN-02", D13), ("C-WN-09", D13 + "; narration scene 6"),
    ("C-FC-01", D11 + "; " + D13 + "; narration scene 8"), ("C-FC-02", D11 + "; " + D13), ("C-FC-03", D11 + "; " + D13),
    ("C-FC-04", D11 + "; " + D13), ("C-FC-08", D13), ("C-FC-09", D13),
    ("C-NS-01", D11 + "; narration scene 8"), ("C-NS-02", D11), ("C-NS-03", D11), ("C-NS-04", D11), ("C-NS-05", D11),
    ("C-53-02", D11), ("C-53-01", D11), ("C-53-03", D11), ("C-53-04", D11),
    ("C-53-06", "narration scene 11"), ("C-53-09", "narration scene 11"), ("C-53-11", "narration scene 11, 16"),
    ("C-53-12", "narration scene 11"), ("C-53-05", "narration metric card"),
    ("C-54-01", D11 + "; " + D12 + "; narration scene 16"), ("C-54-02", D11), ("C-54-03", D11 + "; narration scene 16"),
    ("C-54-04", D12), ("C-54-05", D12 + "; narration scene 11, 16"), ("C-54-06", D12 + "; narration scene 16"),
    ("C-54-07", D12 + "; narration scene 11"),
    ("C-2X-08", D11 + "; narration scene 8"), ("C-2X-09", D11), ("C-2X-10", D11), ("C-2X-11", D11), ("C-2X-07", D11),
    ("C-RG-01", D10 + "; " + D11 + "; " + D14), ("C-RG-02", D11 + "; " + D14),
    ("C-2H-01", "narration scene 16"), ("C-2H-02", "narration scene 16"), ("C-2H-03", "narration scene 16"),
    ("C-2H-07", "narration scene 16"), ("C-46-02", "narration scene 16"),
    ("C-GD-01", D10), ("C-GD-02", D10), ("C-GD-03", D10), ("C-GD-04", D10), ("C-PR-01", D10 + "; narration metric card"),
    ("C-PR-02", D10), ("C-GD-05", D10), ("C-S4-02", D10), ("C-S3-03", D10),
    ("C-S5-06", D13), ("C-S5-01", D15), ("C-S1-01", D15), ("C-PF-02", D14), ("C-TR-01", "narration scene 15"),
    ("C-TR-02", "narration scene 15"), ("C-TR-03", "narration scene 15"), ("C-WN-06", "narration metric card"),
    ("C-FC-20", "narration metric card"), ("C-WN-05", "narration metric card"),
]


# --------------------------------------------------------------------------- expression handling

PH = re.compile(r"\{([A-Za-z0-9_.\-]+)\}")


def py_eval(expr: str) -> float | None:
    vals = {}
    for vid in PH.findall(expr):
        v = V.val(vid)
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            return None
        vals[vid] = v
    code = PH.sub(lambda m: repr(float(vals[m.group(1)])), expr)
    try:
        return float(eval(code, {"__builtins__": {}}, {"ABS": abs, "MIN": min, "MAX": max}))  # noqa: S307
    except (ZeroDivisionError, ValueError):
        return None


def xl_formula(expr: str, ref: dict[str, str]) -> str:
    ids = PH.findall(expr)
    body = PH.sub(lambda m: ref[m.group(1)], expr)
    guard = ",".join(f"ISNUMBER({ref[i]})" for i in dict.fromkeys(ids))
    return f'=IF(AND({guard}),{body},"unavailable")' if ids else f"={body}"


# --------------------------------------------------------------------------- workbook styling

THIN = Side(style="thin", color=LINE)
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BLUE)
NOTE = Font(name="Calibri", color=GREY, italic=True)
TITLE = Font(name="Calibri", bold=True, size=14, color=INK)
LINK = Font(name="Calibri", color=BLUE, underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws: Worksheet, row: int, cols: list[tuple[str, int]]) -> None:
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    ws.row_dimensions[row].height = 32


def finish(ws: Worksheet, head_row: int, ncols: int) -> None:
    last = ws.max_row
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(ncols)}{max(last, head_row + 1)}"
    for row in ws.iter_rows(min_row=head_row + 1, max_row=last, max_col=ncols):
        for c in row:
            c.alignment = WRAP
            c.border = Border(bottom=THIN)
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"  # printable: fit to page width
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"


def title(ws: Worksheet, text: str, note: str) -> None:
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE
    ws["A2"].alignment = Alignment(wrap_text=False, vertical="top")
    back = ws.cell(row=1, column=6, value="← About / navigation")
    back.hyperlink = "#'About'!A1"
    back.font = LINK


def setfmt(cell: Any, fmt: str) -> None:
    cell.number_format = fmt


# --------------------------------------------------------------------------- build

def build() -> dict[str, Any]:
    fusion_values()
    warning_utility_values()
    trend_values()
    other_values()
    deck_values()
    files = {fid: file_row(fid) for fid in FILES}

    wb = Workbook()
    ws_claims = wb.active
    ws_claims.title = "ClaimIndex"
    ws = {n: wb.create_sheet(n) for n in ("Sources", "Cohorts", "Forecasts", "WarningRanges", "Studies", "Gates",
                                          "Verification", "Guardrails", "About")}

    # ---- Sources
    s = ws["Sources"]
    title(s, "Sources — every source value stored once, at full precision",
          "Value column is the only place a source number is typed by the builder; every other sheet references it. "
          "'unavailable' = the file is not on this installation (never 0).")
    cols = [("Value ID", 22), ("Description", 52), ("Value", 16), ("Unit", 14), ("Source ID", 14), ("Repository path", 46),
            ("Locator (JSON pointer / field)", 60), ("Source date", 22), ("Source SHA-256", 30),
            ("Protocol / revision identity", 46), ("Publication status", 24)]
    HR = 4
    header(s, HR, cols)
    ref: dict[str, str] = {}
    for i, r in enumerate(V.rows, HR + 1):
        f = files[r["file"]]
        v = r["value"]
        if v is None:
            v = "unavailable"
        elif isinstance(v, (list, dict)):
            v = json.dumps(v)
        elif isinstance(v, str) and re.fullmatch(r"\d{4}-\d\d-\d\dT[\d:]+(\.\d+)?(\+00:00|Z)?", v):
            v = utc(v)
        row = [r["id"], r["desc"], v, r["unit"], r["file"], f["path"], r["locator"], str(f["date"]), f["sha256"],
               f["identity"], f["status"]]
        for j, x in enumerate(row, 1):
            c = s.cell(row=i, column=j, value=x)
        vc = s.cell(row=i, column=3)
        if isinstance(v, float):
            vc.number_format = "0.0000"
        elif isinstance(v, dt.datetime):
            vc.number_format = "yyyy-mm-dd hh:mm:ss"
        ref[r["id"]] = f"Sources!$C${i}"
    finish(s, HR, len(cols))
    # file index to the right is avoided (one table per sheet); the JSON export carries the per-file list.

    # ---- Cohorts
    co = ws["Cohorts"]
    title(co, "Cohorts — who each result counts", "Separate evaluation groups; never compare rows across groups.")
    cols = [("Evaluation group", 30), ("Dataset / version", 26), ("Participants", 12), ("What the participant count means", 36),
            ("Cadence", 12), ("Split roles", 40), ("Horizon (min)", 12), ("Origins / forecasts", 14), ("Hours", 10),
            ("Onsets", 10), ("Inspected / unseen status", 40), ("Opening disclosure", 36), ("Model scope", 34), ("Source IDs", 30)]
    header(co, 4, cols)
    cohorts = [
        ["Serving test (final test segment)", "CGMacros 1.0.0", "{fus.participants}", "Type 2 group (HbA1c > 6.4 %), same people as training",
         "5 min", "Chronological 60/10/10/20 within each person: training / selection / calibration / test, purge gaps", "30, 60",
         "{fus.seg.test.origins}", "{fus.seg.test.covered_hours}", "{fus.seg.test.distinct_new_onsets_60}",
         "Inspected: the v1/v2 test segment was read before the fusion constraint was declared (retrospective)",
         "Later recordings of the same 14 people — never 'unseen patients'", "Serving fusion LightGBM + correction",
         "fus.participants; fus.seg.test.*"],
        ["Serving model selection (validation segment)", "CGMacros 1.0.0", "{fus.participants}", "Same 14 people", "5 min",
         "Validation segment used for model, threshold and adaptation choice", "30, 60", "{fus.seg.validation.origins}",
         "{fus.seg.validation.covered_hours}", "{fus.seg.validation.distinct_new_onsets_60}", "Used for selection", "—",
         "Serving candidates", "fus.seg.validation.*"],
        ["Serving calibration segment", "CGMacros 1.0.0", "{fus.participants}", "Same 14 people", "5 min", "Split-conformal range calibration",
         "30, 60", "{fus.seg.calibration.origins}", "{fus.seg.calibration.covered_hours}", "{fus.seg.calibration.distinct_new_onsets_60}",
         "Used for calibration", "—", "Serving ranges", "fus.seg.calibration.*"],
        ["Excluded CGMacros participants", "CGMacros 1.0.0", "{fus.excluded}", "Pre-diabetes or no diabetes; replay only", "5 min",
         "Never used in training or evaluation", "—", "—", "—", "—", "Not evaluated", "—", "Replay only", "fus.excluded; inv.cgm"],
        ["Participant-held-out internal (D-032 nested; D-043…D-053 reuse the folds)", "CGMacros 1.0.0", "{inv.cgm_eval}",
         "Each of the 14 people held out entirely once (5 outer folds)", "5 min", "Outer folds assess; inner folds select; calibration people for ranges",
         "30, 60 (120 in D-043)", "—", "—", "—", "Previously inspected people: internal validation, not an untouched cohort",
         "Each study opens its outer folds once behind a lock file", "Research procedures (not served)", "inv.cgm_eval; nest.*"],
        ["External (D-032 first opening; D-045 second opening)", "ShanghaiT2DM, figshare 21600933 v5", "{inv.sh_eval}",
         "Type 2 people imported (100) with usable recordings", "15 min", "Assessment only; nothing chosen on this cohort", "30, 60, 120",
         "—", "—", "—", "Reserved external cohort", "Opened twice (D-032 30/60; D-045 120)", "CGM-only 15-min LightGBM — not the fusion model",
         "inv.sh; inv.sh_eval"],
        ["Foundation baseline (D-054)", "CGMacros 1.0.0", "{fus.participants}", "All 14 Type 2 people, all eligible origins", "5 min",
         "No training (zero-shot)", "30, 60", "{fd.origins}", "—", "—", "Previously inspected people", "—", "Chronos-Bolt-Tiny, CGM only",
         "fd.origins"],
        ["Synthetic demonstration", "Fictional synthetic participants", "—", "Fictional; labelled synthetic everywhere", "5 min", "—", "30, 60",
         "—", "—", "—", "Never evidence of performance", "—", "synthetic-demo.lightgbm.dyn-v2.v1", "reg.synthetic; reg.synthetic_fitted"],
    ]
    for i, r in enumerate(cohorts, 5):
        for j, x in enumerate(r, 1):
            if isinstance(x, str) and PH.fullmatch(x):
                x = xl_formula(x, ref)
            co.cell(row=i, column=j, value=x)
        for j in (3, 8, 9, 10):
            co.cell(row=i, column=j).number_format = "#,##0"
    finish(co, 4, len(cols))

    # ---- Forecasts
    fo = ws["Forecasts"]
    title(fo, "Forecasts — errors by evaluation, model and horizon",
          "MAE = mean absolute error, macro (equal weight per participant). Difference = candidate − persistence on the same "
          "origins (negative = lower error). Saved paired differences and CIs are participant-bootstrap values from the artifacts.")
    cols = [("Evaluation group", 30), ("Model / procedure", 40), ("Horizon (min)", 9), ("Aggregation", 16), ("MAE (mg/dL)", 11),
            ("Persistence MAE, same origins", 13), ("Candidate − persistence (formula)", 14), ("Relative change % (formula)", 12),
            ("RMSE (mg/dL)", 11), ("Saved paired diff", 11), ("CI 95 % low", 10), ("CI 95 % high", 10), ("Note", 40), ("Source IDs", 36)]
    header(fo, 4, cols)
    frows = []
    for key, name in (("serv", "Fusion LightGBM + correction (SERVING)"), ("cbpa", "CatBoost + profile + correction (not selected)"),
                      ("lgbp", "LightGBM + profile, population"), ("v2a", "LightGBM dynamic-only + correction (v2)"),
                      ("ridgep", "Ridge dynamic + profile"), ("trend", "Linear trend"), ("pers", "Persistence")):
        for h in (30, 60):
            sp = key == "serv"
            frows.append([SERV, name, h, "macro per participant", f"{{fus.{key}.mae{h}}}", f"{{fus.pers.mae{h}}}",
                          f"{{fus.{key}.rmse{h}}}", f"{{fus.cmp.sp.d{h}}}" if sp else None, f"{{fus.cmp.sp.lo{h}}}" if sp else None,
                          f"{{fus.cmp.sp.hi{h}}}" if sp else None, "Retrospective; same 14 people as training" if sp else ""])
    for key, name in (("sel", "Selected research procedure (Ridge CGM-only in 4/5 folds)"), ("v2", "v2 recipe refit per fold")):
        for h in (30, 60):
            frows.append([NEST + " (D-032)", name, h, "macro per participant", f"{{nest.{key}.mae{h}}}", f"{{nest.pers.mae{h}}}",
                          None, None, None, None, "Not served"])
    for h in (30, 60):
        frows.append([EXT + " (D-032)", "15-min LightGBM, CGM only", h, "macro per participant", f"{{ext.cand.mae{h}}}",
                      f"{{ext.pers.mae{h}}}", None, f"{{ext.cmp.mean_diff{h}}}", f"{{ext.cmp.lo{h}}}", f"{{ext.cmp.hi{h}}}",
                      "Not the fusion model"])
    for h in (30, 60, 120):
        frows.append([NEST + " (D-043)", "Two-hour selected procedure, personalised", h, "macro per participant", f"{{2h.sel.mae{h}}}",
                      f"{{2h.pers.mae{h}}}", None, "{2h.cmp.d}" if h == 120 else None, "{2h.cmp.lo}" if h == 120 else None,
                      "{2h.cmp.hi}" if h == 120 else None, "Not served"])
        frows.append([EXT + " (D-045)", "15-min LightGBM refit for 120 min", h, "macro per participant", f"{{2hx.mae{h}}}",
                      f"{{2hx.pers{h}}}", None, "{2hx.cmp.mean_diff}" if h == 120 else None, "{2hx.cmp.lo}" if h == 120 else None,
                      "{2hx.cmp.hi}" if h == 120 else None, "Second opening of the cohort"])
    for h in (30, 60):
        frows.append([NEST + " (D-053)", "Constrained selection S (not adopted)", h, "macro per participant", f"{{cs.S.mae{h}}}",
                      f"{{cs.persistence.mae{h}}}", None, None, None, None, "Gate failed; not adopted"])
        frows.append([NEST + " (D-053)", "Serving recipe refitted in the same folds (R0)", h, "macro per participant", f"{{cs.R0.mae{h}}}",
                      f"{{cs.persistence.mae{h}}}", None, "{cs.p60.mean_diff}" if h == 60 else None, "{cs.p60.lo}" if h == 60 else None,
                      "{cs.p60.hi}" if h == 60 else None, "Saved diff on this row = S − R0" if h == 60 else ""])
        frows.append([NEST + " — all origins (D-054)", "Chronos-Bolt-Tiny zero-shot, CGM only", h, "macro per participant", f"{{fd.FM{h}}}",
                      f"{{fd.P{h}}}", None, "{fd.fp.mean}" if h == 60 else None, "{fd.fp.lo}" if h == 60 else None,
                      "{fd.fp.hi}" if h == 60 else None, "Not fine-tuned, not served"])
    fstart = 5
    for i, r in enumerate(frows, fstart):
        grp, name, h, agg, mae, pmae, rmse, d, lo, hi, note = r
        ids = [x for x in (mae, pmae, rmse, d, lo, hi) if x]
        fo.cell(row=i, column=1, value=grp)
        fo.cell(row=i, column=2, value=name)
        fo.cell(row=i, column=3, value=h)
        fo.cell(row=i, column=4, value=agg)
        fo.cell(row=i, column=5, value=xl_formula(mae, ref))
        fo.cell(row=i, column=6, value=xl_formula(pmae, ref))
        fo.cell(row=i, column=7, value=f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(F{i})),E{i}-F{i},"unavailable")')
        fo.cell(row=i, column=8, value=f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(F{i})),(E{i}-F{i})/F{i}*100,"unavailable")')
        fo.cell(row=i, column=9, value=xl_formula(rmse, ref) if rmse else "not in source")
        for j, x in ((10, d), (11, lo), (12, hi)):
            fo.cell(row=i, column=j, value=xl_formula(x, ref) if x else "—")
        fo.cell(row=i, column=13, value=note)
        fo.cell(row=i, column=14, value="; ".join(PH.findall(" ".join(ids))))
        for j in (5, 6, 7, 9, 10, 11, 12):
            fo.cell(row=i, column=j).number_format = "0.00"
        fo.cell(row=i, column=8).number_format = "0.0"
    fend = fstart + len(frows) - 1
    finish(fo, 4, len(cols))
    # chart data block (references the table) + chart 1
    cd = fend + 3
    fo.cell(row=cd, column=1, value="Chart data: serving test MAE vs persistence (references the table above)").font = NOTE
    fo.cell(row=cd + 1, column=1, value="Horizon")
    fo.cell(row=cd + 1, column=2, value="Serving fusion LightGBM")
    fo.cell(row=cd + 1, column=3, value="Persistence")
    for k, h in enumerate((30, 60)):
        rr = fstart + k  # serving rows are the first two
        fo.cell(row=cd + 2 + k, column=1, value=f"{h} min")
        fo.cell(row=cd + 2 + k, column=2, value=f"=E{rr}")
        fo.cell(row=cd + 2 + k, column=3, value=f"=F{rr}")
        fo.cell(row=cd + 2 + k, column=2).number_format = "0.00"
        fo.cell(row=cd + 2 + k, column=3).number_format = "0.00"
    ch = BarChart()
    ch.type = "col"
    ch.title = "Serving test MAE vs persistence (mg/dL, lower is better)"
    ch.y_axis.title = "MAE (mg/dL)"
    ch.add_data(Reference(fo, min_col=2, max_col=3, min_row=cd + 1, max_row=cd + 3), titles_from_data=True)
    ch.set_categories(Reference(fo, min_col=1, min_row=cd + 2, max_row=cd + 3))
    ch.series[0].graphicalProperties.solidFill = BLUE
    ch.series[1].graphicalProperties.solidFill = "9CA3AF"
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.height, ch.width = 8, 14
    fo.add_chart(ch, f"E{cd}")

    # ---- WarningRanges
    wr = ws["WarningRanges"]
    title(wr, "Warnings and prediction ranges",
          "Numerator and denominator are shown for every share. Notice shares use ALL onsets; the median lead uses DETECTED onsets. "
          "False-alarm rates divide by the saved observed / eligible patient-day denominators (not by rounded hours).")
    cols = [("Section", 22), ("Metric", 46), ("Value (formula)", 13), ("Unit", 12), ("Numerator", 12), ("Denominator", 12),
            ("Definition / denominator meaning", 56), ("Saved value to agree with", 13), ("Agrees?", 10), ("Source IDs", 34)]
    header(wr, 4, cols)
    wrows = [
        ("Serving warning (test)", "Onsets (eligible new sustained > 180 episodes)", "{fus.ev.true_onsets}", "onsets", None, None, "Distinct new onsets; not origins", None),
        ("Serving warning (test)", "Detected onsets", "{fus.ev.detected_onsets}", "onsets", None, None, "Onsets matched to a warning episode", None),
        ("Serving warning (test)", "Missed onsets", "{fus.ev.true_onsets}-{fus.ev.detected_onsets}", "onsets", None, None, "", None),
        ("Serving warning (test)", "Detection share (detected / all onsets)", "{fus.ev.detected_onsets}/{fus.ev.true_onsets}", "share",
         "{fus.ev.detected_onsets}", "{fus.ev.true_onsets}", "Recall over all onsets", None),
        ("Serving warning (test)", "Warned ≥ 15 min ahead, share of ALL onsets", "{wu.ge15}/{wu.onsets}", "share", "{wu.ge15}", "{wu.onsets}",
         "All onsets; a missed onset counts against it", None),
        ("Serving warning (test)", "Warned ≥ 30 min ahead, share of ALL onsets", "{wu.ge30}/{wu.onsets}", "share", "{wu.ge30}", "{wu.onsets}",
         "All onsets (headline 48.4 %)", None),
        ("Serving warning (test)", "Warned ≥ 60 min ahead, share of ALL onsets", "{wu.ge60}/{wu.onsets}", "share", "{wu.ge60}", "{wu.onsets}", "All onsets", None),
        ("Serving warning (test)", "Median qualified lead (DETECTED onsets only)", "{fus.ev.qualified_lead_median}", "min", None,
         "{fus.ev.qualified_lead_n}", "Denominator = detected onsets with a qualified lead; not comparable to the all-onset shares", "{wu.lead_median}"),
        ("Serving warning (test)", "Qualified lead IQR low (detected onsets)", "{fus.ev.iqr_lo}", "min", None, None, "", "{wu.lead_q25}"),
        ("Serving warning (test)", "Qualified lead IQR high (detected onsets)", "{fus.ev.iqr_hi}", "min", None, None, "", "{wu.lead_q75}"),
        ("Serving warning (test)", "False-alarm episodes", "{fus.ev.false_alarm_episodes}", "episodes", None, None, "", None),
        ("Serving warning (test)", "Warning episodes", "{fus.ev.warning_episodes}", "episodes", None, None, "", None),
        ("Serving warning (test)", "Precision (true episodes / episodes)", "({fus.ev.warning_episodes}-{fus.ev.false_alarm_episodes})/{fus.ev.warning_episodes}",
         "ratio", None, "{fus.ev.warning_episodes}", "", "{fus.ev.precision}"),
        ("Serving warning (test)", "False alarms per OBSERVED patient-day", "{fus.ev.false_alarm_episodes}/{fus.ev.observed_patient_days}",
         "per day", "{fus.ev.false_alarm_episodes}", "{fus.ev.observed_patient_days}",
         "Observed patient-days from saved event metadata (time base: issued origins × 5 min)", "{fus.ev.false_alarms_per_observed_patient_day}"),
        ("Serving warning (test)", "False alarms per ELIGIBLE patient-day", "{fus.ev.false_alarm_episodes}/{fus.ev.eligible_patient_days}",
         "per day", "{fus.ev.false_alarm_episodes}", "{fus.ev.eligible_patient_days}", "Eligible patient-days (below-threshold origins) — a different, smaller base",
         "{fus.ev.false_alarms_per_patient_day}"),
        ("Serving warning (test)", "Time in warning (share of eligible origins)", "{wu.firing}/{wu.eligible}", "share", "{wu.firing}", "{wu.eligible}",
         "Post-hoc diagnostic", "{wu.share_eligible}"),
        ("Serving warning (test)", "Warning margin (score = max forecast − 180)", "{fus.ev.threshold}", "mg/dL", None, None, "Chosen on validation by event F1", None),
        ("Persistence warning (test)", "False alarms per observed patient-day", "{fus.pev.false_alarms_per_observed_patient_day}", "per day", None, None, "Fires at 150 mg/dL", None),
        ("Serving ranges (test)", "80 % range coverage, 30 min", "{fus.iv.coverage30}", "share", None, "{fus.iv.n30}", "Pooled over forecasts", None),
        ("Serving ranges (test)", "80 % range coverage, 60 min", "{fus.iv.coverage60}", "share", None, "{fus.iv.n60}", "Pooled over forecasts", None),
        ("Serving ranges (test)", "Coverage gap vs target, 30 min (points)", "({fus.iv.coverage30}-{const.target_cov})*100", "points", None, None, "Negative = under-coverage", None),
        ("Serving ranges (test)", "Coverage gap vs target, 60 min (points)", "({fus.iv.coverage60}-{const.target_cov})*100", "points", None, None, "Negative = under-coverage", None),
        ("Serving ranges (test)", "Per-person coverage min, 30 min", "{fus.iv.per_patient_min30}", "share", None, None, "Weakest participant", None),
        ("Serving ranges (test)", "Per-person coverage max, 30 min", "{fus.iv.per_patient_max30}", "share", None, None, "", None),
        ("Serving ranges (test)", "Per-person coverage min, 60 min", "{fus.iv.per_patient_min60}", "share", None, None, "Weakest participant", None),
        ("Serving ranges (test)", "Per-person coverage max, 60 min", "{fus.iv.per_patient_max60}", "share", None, None, "", None),
        ("Serving ranges (test)", "Range width (full), 30 min", "{fus.iv.width30}", "mg/dL", None, None, "Full width: 2 × half-width q", None),
        ("Serving ranges (test)", "Range width (full), 60 min", "{fus.iv.width60}", "mg/dL", None, None, "Full width: 2 × half-width q", None),
    ]
    for t in ("falling_fast", "falling", "steady", "rising", "rising_fast"):
        for h in (30, 60):
            wrows.append(("Trend context — VALIDATION", f"{t}, {h} min: macro MAE", f"{{tr.{t}.mae{h}}}", "mg/dL", None, f"{{tr.{t}.n{h}}}",
                          "Validation segment (selection data), not the test", None))
            wrows.append(("Trend context — VALIDATION", f"{t}, {h} min: served range coverage", f"{{tr.{t}.cov{h}}}", "share", None,
                          f"{{tr.{t}.n{h}}}", "Validation segment", None))
    for st in ("rising", "steady", "falling"):
        for h in (30, 60):
            wrows.append(("Trend context — FINAL TEST (D-050, read once)", f"{st}, {h} min: served range coverage", f"{{rg.inc.{st}{h}}}", "share",
                          None, f"{{rg.inc.{st}n{h}}}", "Final test, retrospective; rising/falling = |slope| > 1 mg/dL/min", None))
    for i, (sec, met, e, unit, num, den, dfn, saved) in enumerate(wrows, 5):
        wr.cell(row=i, column=1, value=sec)
        wr.cell(row=i, column=2, value=met)
        wr.cell(row=i, column=3, value=xl_formula(e, ref))
        wr.cell(row=i, column=4, value=unit)
        wr.cell(row=i, column=5, value=xl_formula(num, ref) if num else "—")
        wr.cell(row=i, column=6, value=xl_formula(den, ref) if den else "—")
        wr.cell(row=i, column=7, value=dfn)
        if saved:
            wr.cell(row=i, column=8, value=xl_formula(saved, ref))
            wr.cell(row=i, column=9, value=f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(H{i})),IF(ABS(C{i}-H{i})<1E-9*MAX(1,ABS(H{i})),"agrees","DIFFERS"),"n/a")')
            wr.cell(row=i, column=8).number_format = "0.0000"
        else:
            wr.cell(row=i, column=8, value="—")
            wr.cell(row=i, column=9, value="—")
        wr.cell(row=i, column=10, value="; ".join(dict.fromkeys(PH.findall(" ".join(x for x in (e, num, den, saved) if x)))))
        wr.cell(row=i, column=3).number_format = "0.0%" if unit == "share" else ("0.00" if unit not in ("onsets", "episodes") else "0")
        for j in (5, 6):
            wr.cell(row=i, column=j).number_format = "#,##0.##"
    wend = 4 + len(wrows)
    finish(wr, 4, len(cols))
    # chart 2: range coverage vs the 80 % target
    cd = wend + 3
    wr.cell(row=cd, column=1, value="Chart data: 80 % range coverage by evaluation (target line = 80 %)").font = NOTE
    for j, n in enumerate(("Evaluation", "Coverage", "Target"), 1):
        wr.cell(row=cd + 1, column=j, value=n)
    cov_rows = [("Serving 30 min (test)", "{fus.iv.coverage30}"), ("Serving 60 min (test)", "{fus.iv.coverage60}"),
                ("Nested research 30 min", "{nest.cov30}"), ("Nested research 60 min", "{nest.cov60}"),
                ("D-050 candidate 30 (not served)", "{rg.pool.cov30}"), ("D-050 candidate 60 (not served)", "{rg.pool.cov60}"),
                ("D-053 S 30 (not adopted)", "{cs.cov30}"), ("D-053 S 60 (not adopted)", "{cs.cov60}"),
                ("Two-hour 120 (D-043)", "{2h.cov120}"), ("Two-hour pooled 120 (D-046)", "{wr.b1.cov120}"),
                ("External 120 (D-045)", "{2hx.cov120}")]
    for k, (lab, e) in enumerate(cov_rows):
        r_ = cd + 2 + k
        wr.cell(row=r_, column=1, value=lab)
        wr.cell(row=r_, column=2, value=xl_formula(e, ref)).number_format = "0.0%"
        wr.cell(row=r_, column=3, value=f"={ref['const.target_cov']}").number_format = "0.0%"
    ch = BarChart()
    ch.type = "bar"
    ch.title = "80 % prediction-range coverage vs the 80 % target"
    ch.x_axis.title = None
    ch.y_axis.title = "Coverage"
    ch.y_axis.scaling.min = 0.6
    ch.y_axis.scaling.max = 0.9
    ch.y_axis.number_format = "0%"
    ch.add_data(Reference(wr, min_col=2, max_col=3, min_row=cd + 1, max_row=cd + 1 + len(cov_rows)), titles_from_data=True)
    ch.set_categories(Reference(wr, min_col=1, min_row=cd + 2, max_row=cd + 1 + len(cov_rows)))
    ch.series[0].graphicalProperties.solidFill = BLUE
    ch.series[1].graphicalProperties.solidFill = "D1D5DB"
    ch.height, ch.width = 10, 16
    wr.add_chart(ch, f"E{cd}")

    # ---- Gates
    ga = ws["Gates"]
    title(ga, "Gates — the declared rules applied to unrounded saved values",
          "Value column holds the unrounded source value (display format may round). Result = formula; 'Saved outcome' = the "
          "artifact's own flag. A study passes only if every criterion row passes.")
    cols = [("Study", 10), ("Gate group", 18), ("Criterion (as declared)", 50), ("Metric value (unrounded)", 14),
            ("Comparator", 10), ("Limit", 10), ("Result (formula)", 12), ("Saved outcome", 12), ("Agrees with saved?", 12),
            ("Sign convention / note", 50), ("Source IDs", 34)]
    header(ga, 4, cols)
    G: list[tuple] = [
        # study, group, criterion, expr, comparator, limit, saved expr (TRUE = criterion passed), note
        ("D-048", "Development target", "Selected alpha improves 60-min dev MAE vs alpha = 1 by ≥ 0.5 mg/dL", "{blend.improvement_60}", ">=", 0.5,
         None, "Improvement positive = gain (alpha = 1 minus selected)"),
        ("D-048", "Guardrail", "30-min worsening ≤ 0.25", "{blend.worsening_30}", "<=", 0.25, None, "Candidate − control"),
        ("D-048", "Guardrail", "High-glucose 60-min worsening ≤ 1.0", "{blend.high_worsening_60}", "<=", 1.0, None, "Candidate − control"),
        ("D-048", "Guardrail", "Availability drop ≤ 0.5 points", "{blend.availability_drop_points}", "<=", 0.5, None, ""),
        ("D-049", "Gate (A1)", "A1 60-min personalised gain vs B0 ≥ 0.5", "{rt.A1.gain_60}", ">=", 0.5, None, "Gain positive = lower error (B0 − arm)"),
        ("D-049", "Gate (A2)", "A2 60-min personalised gain vs B0 ≥ 0.5", "{rt.A2.gain_60}", ">=", 0.5, None, "Gain positive = lower error"),
        ("D-049", "Gate (B1)", "B1 60-min personalised gain vs B0 ≥ 0.5", "{rt.B1.gain_60}", ">=", 0.5, None, "Gain positive = lower error"),
        ("D-050", "Adoption (T replaces G)", "Pre-test mean |coverage − 0.80|: T better than G by ≥ 1.0 point",
         "({rg.err.G}-{rg.err.T})*100", ">=", 1.0, None, "G − T in points; positive = T closer to 80 %"),
        ("D-050", "Adoption (T replaces G)", "T mean width ≤ 1.15 × G width", "{rg.width.T}/{rg.width.G}", "<=", 1.15, None, "Ratio"),
        ("D-051", "Adoption", "C1 − C0 at 60 min ≤ −0.5 mg/dL", "{tu.c1c0.mean_diff}", "<=", -0.5, None, "Candidate − control (negative = gain)"),
        ("D-051", "Adoption", "Bootstrap 95 % CI upper < 0", "{tu.c1c0.hi}", "<", 0, None, "Participant bootstrap"),
        ("D-051", "Adoption", "30-min worsening ≤ 0.25", "{tu.adopt.worsening_30}", "<=", 0.25, None, "Candidate − control"),
        ("D-051", "Adoption", "High-glucose 60-min worsening ≤ 1.0", "{tu.adopt.high_worsening_60}", "<=", 1.0, None, "Candidate − control"),
        ("D-052", "Primary", "S − R0 at 60 min ≤ −0.5", "{jt.gate.mean_diff_60}", "<=", -0.5, None, "Candidate − control"),
        ("D-052", "Primary", "95 % CI upper < 0", "{jt.gate.ci95_60_hi}", "<", 0, None, ""),
        ("D-052", "Primary", "30-min worsening ≤ 0.25", "{jt.gate.worsening_30}", "<=", 0.25, None, ""),
        ("D-052", "Primary", "High-target (> 180) 60-min worsening ≤ 1.0", "{jt.gate.high_worsening_60}", "<=", 1.0, None, "Fails"),
        ("D-052", "Primary", "Availability drop ≤ 0.5 points", "{jt.gate.availability_drop_points}", "<=", 0.5, None, ""),
        ("D-052", "Release guardrail", "False alarms/observed day ≤ R0 + 0.10", "{jt.rel.fa_per_day_delta}", "<=", 0.10, None, "S − R0"),
        ("D-052", "Release guardrail", "Time in warning ≤ R0 + 2.0 points", "{jt.rel.time_in_warning_delta_points}", "<=", 2.0, None, "S − R0"),
        ("D-052", "Release guardrail", "Share of all onsets warned ≥ 30 min ≥ R0 − 2.0 points", "{jt.rel.warned30_delta_points}", ">=", -2.0, None, "S − R0"),
        ("D-053", "Primary", "S − R0 at 60 min ≤ −0.5 (inclusive)", "{cs.pg.mean_diff_60}", "<=", -0.5, "{cs.pgok.mean_diff_60}", "Candidate − control"),
        ("D-053", "Primary", "Paired bootstrap 95 % CI upper < 0 (strict)", "{cs.pg.ci95_60_hi}", "<", 0, "{cs.pgok.ci95_upper_below_0}", "2,000 resamples"),
        ("D-053", "Primary", "30-min worsening ≤ +0.25 (inclusive)", "{cs.pg.worsening_30}", "<=", 0.25, "{cs.pgok.worsening_30}", ""),
        ("D-053", "Primary", "High-target (> 180) 60-min worsening ≤ +1.0 (inclusive)", "{cs.pg.high_worsening_60}", "<=", 1.0,
         "{cs.pgok.high_worsening_60}", "Unrounded ≈ +1.0658 > +1.0 → FAIL at any display rounding"),
        ("D-053", "Primary", "Availability drop ≤ 0.5 points", "{cs.pg.availability_drop_points}", "<=", 0.5, "{cs.pgok.availability_drop}", ""),
        ("D-053", "Operational", "False alarms/observed day ≤ R0-P + 0.10", "{cs.w.S.fa_day}-{cs.w.R0P.fa_day}", "<=", 0.10, None, "S − R0-P"),
        ("D-053", "Operational", "Time in warning ≤ R0-P + 2.0 points", "({cs.w.S.tiw}-{cs.w.R0P.tiw})*100", "<=", 2.0, None, "S − R0-P"),
        ("D-053", "Operational", "Share of all onsets warned ≥ 30 min ≥ R0-P − 2.0 points", "({cs.w.S.share30}-{cs.w.R0P.share30})*100", ">=", -2.0, None, "S − R0-P"),
    ]
    op = {"<=": "<=", ">=": ">=", "<": "<", ">": ">"}
    grow: dict[str, list[int]] = {}
    for i, (study, grp, crit, e, cmp_, lim, saved, note) in enumerate(G, 5):
        ga.cell(row=i, column=1, value=study)
        ga.cell(row=i, column=2, value=grp)
        ga.cell(row=i, column=3, value=crit)
        ga.cell(row=i, column=4, value=xl_formula(e, ref)).number_format = "0.0000"
        ga.cell(row=i, column=5, value=cmp_)
        ga.cell(row=i, column=6, value=lim)
        ga.cell(row=i, column=7, value=f'=IF(ISNUMBER(D{i}),IF(D{i}{op[cmp_]}F{i},"PASS","FAIL"),"FAIL (undefined)")')
        if saved:
            ga.cell(row=i, column=8, value=f"={ref[PH.findall(saved)[0]]}")
            ga.cell(row=i, column=9, value=f'=IF(OR(H{i}="unavailable",H{i}=""),"n/a",IF((G{i}="PASS")=(UPPER(H{i})="TRUE"),"agrees","DIFFERS"))')
        else:
            ga.cell(row=i, column=8, value="—")
            ga.cell(row=i, column=9, value="—")
        ga.cell(row=i, column=10, value=note)
        ga.cell(row=i, column=11, value="; ".join(PH.findall(e) + (PH.findall(saved) if saved else [])))
        grow.setdefault(study, []).append(i)
    gl = 5 + len(G)
    decisions = {"D-048": "{blend.target_met}", "D-049": None, "D-050": "{rg.chosen}", "D-051": "{tu.adopt.c1_better}",
                 "D-052": "{jt.gate.passed}", "D-053": "{cs.all_passed}"}
    gate_cells: dict[str, str] = {}
    for k, (study, saved) in enumerate(decisions.items()):
        i = gl + 1 + k
        rows_ = grow[study]
        rng = f"G{rows_[0]}:G{rows_[-1]}"
        ga.cell(row=i, column=1, value=study).font = Font(bold=True, color=INK)
        ga.cell(row=i, column=2, value="OVERALL")
        if study == "D-049":
            ga.cell(row=i, column=3, value="Any arm passes its gate → that arm qualifies")
            ga.cell(row=i, column=7, value=f'=IF(COUNTIF({rng},"PASS")>0,"PASSED","FAILED")')
        elif study == "D-050":
            ga.cell(row=i, column=3, value="Both adoption criteria → T adopted; else G kept")
            ga.cell(row=i, column=7, value=f'=IF(COUNTIF({rng},"FAIL*")=0,"PASSED","FAILED")')
        else:
            ga.cell(row=i, column=3, value="All criteria above pass")
            ga.cell(row=i, column=7, value=f'=IF(COUNTIF({rng},"FAIL*")=0,"PASSED","FAILED")')
        ga.cell(row=i, column=7).font = Font(bold=True, color=BLUE)
        if saved:
            ga.cell(row=i, column=8, value=f"={ref[PH.findall(saved)[0]]}")
        ga.cell(row=i, column=10, value={"D-048": "Not met → incumbent retained; outer folds not opened",
                                         "D-049": "No arm met the gate; B0 retained; outer folds not opened",
                                         "D-050": "Saved choice G (served shape kept)",
                                         "D-051": "Not adopted (small internal gain vs equally tuned C0 is not a release)",
                                         "D-052": "Not adopted; serving unchanged",
                                         "D-053": "Primary gate FAILED (high-target); operational passed; NOT ADOPTED"}[study])
        gate_cells[study] = f"Gates!$G${i}"
    # D-053 sub-gate summaries for the operational claim
    i = gl + 1 + len(decisions)
    ga.cell(row=i, column=1, value="D-053")
    ga.cell(row=i, column=2, value="OPERATIONAL")
    ga.cell(row=i, column=3, value="Operational criteria pass")
    op_rows = [r for r in grow["D-053"] if ga.cell(row=r, column=2).value == "Operational"]
    ga.cell(row=i, column=7, value=f'=IF(COUNTIF(G{op_rows[0]}:G{op_rows[-1]},"FAIL*")=0,"PASSED","FAILED")')
    ga.cell(row=i, column=8, value=f"={ref['cs.op.passed']}")
    gate_cells["D-053 operational"] = f"Gates!$G${i}"
    i += 1
    ga.cell(row=i, column=1, value="D-053")
    ga.cell(row=i, column=2, value="PRIMARY")
    ga.cell(row=i, column=3, value="Primary point criteria pass")
    pr_rows = [r for r in grow["D-053"] if ga.cell(row=r, column=2).value == "Primary"]
    ga.cell(row=i, column=7, value=f'=IF(COUNTIF(G{pr_rows[0]}:G{pr_rows[-1]},"FAIL*")=0,"PASSED","FAILED")')
    ga.cell(row=i, column=8, value=f"={ref['cs.pg.passed']}")
    gate_cells["D-053 primary"] = f"Gates!$G${i}"
    i += 1
    ga.cell(row=i, column=1, value="D-054")
    ga.cell(row=i, column=2, value="—")
    ga.cell(row=i, column=3, value="No adoption gate: descriptive research baseline (F1 residual reported, not selected)")
    ga.cell(row=i, column=7, value="not applicable")
    finish(ga, 4, len(cols))

    # Python-side gate outcomes (for the JSON export and claim checks)
    def gate_py(study: str, group: str | None = None) -> str:
        res = []
        for st, grp, _c, e, cmp_, lim, _s, _n in G:
            if st != study or (group and grp != group):
                continue
            v = py_eval(e)
            ok = v is not None and {"<=": v <= lim, ">=": v >= lim, "<": v < lim, ">": v > lim}[cmp_]
            res.append(ok)
        if study == "D-049" and not group:
            return "PASSED" if any(res) else "FAILED"
        return "PASSED" if res and all(res) else "FAILED"

    gate_py_out = {s_: gate_py(s_) for s_ in decisions} | {"D-053 operational": gate_py("D-053", "Operational"),
                                                           "D-053 primary": gate_py("D-053", "Primary")}

    # ---- Studies
    st = ws["Studies"]
    title(st, "Studies — D-043 onward, including every negative or unopened result",
          "Protocol-before-run = protocol file's first git commit precedes the saved run/open time (git dates are local "
          "commit evidence, not a clinical registration). 'not opened' / 'unavailable' are never zero.")
    cols = [("Decision", 9), ("Experiment ID", 24), ("Kind", 22), ("Hypothesis / purpose", 44), ("Population and role", 34),
            ("Protocol file", 34), ("Protocol first commit (UTC)", 18), ("Run start / assessment opened (UTC)", 18),
            ("Protocol-before-run (formula)", 16), ("Fit budget", 18), ("Assessment-opening status", 26), ("Candidate metric", 12),
            ("Control metric", 12), ("Candidate − control (formula)", 13), ("Gate outcome (Gates sheet)", 14), ("Adoption decision", 30),
            ("Negative findings / limits", 44), ("Report", 30), ("Artifact status", 24)]
    header(st, 4, cols)

    def commit_utc(proto: str) -> dt.datetime | None:
        _h, d = first_commit(proto)
        return utc(d) if d else None

    studies = [
        ("D-032", "gt-nested-2026-10-12-v1", "Nested internal + external", "New-participant and external evidence for the procedure",
         "14 CGMacros people held out once; 98 ShanghaiT2DM external", "{nest.opened_at}", "—", "Opened once (internal and external)",
         "{nest.sel.mae60}", "{nest.pers.mae60}", None, "Research candidate, not promoted",
         "Git commit of the protocol file postdates its freeze/open timestamps; the frozen protocol hash in report.json is the evidence",
         "reports/nested_benchmark_report.md", "nested"),
        ("D-043", "gt-2h-2026-10-20", "Two-hour forecast + warning", "Does a 120-min forecast and warning work on people held out?",
         "14 CGMacros people held out once (inspected)", "{2h.opened_at}", "Same procedure as D-032", "Opened once",
         "{2h.sel.mae120}", "{2h.pers.mae120}", None, "Research only; demo model unchanged",
         "Median lead 15 min; ranges under-cover (76.5 %)", "reports/two_hour_benchmark.md", "twoh"),
        ("D-044", "gt-2h-clf-2026-10-20", "Direct onset classifier", "Does a classifier trained on the event warn earlier at the same workload?",
         "Folds of D-043", "{clf.opened_at}", "Fixed settings", "Opened, then re-opened once (post-hoc diagnostic, reason logged)",
         None, None, None, "Rejected (negative result)", "Threshold hit the grid minimum; warning active ~99 % of eligible time",
         "reports/two_hour_classifier.md", "clf"),
        ("D-045", "gt-2h-ext-2026-10-20", "External two-hour", "Does the 120-min forecast hold on an independent cohort?",
         "98 ShanghaiT2DM people (external)", "{2hx.opened_at}", "15-min track of D-032", "Second opening of the external cohort",
         "{2hx.mae120}", "{2hx.pers120}", None, "Research only", "CGM only; not the fusion model", "reports/two_hour_external.md", "twoh_ext"),
        ("D-046", "gt-warn-range-2026-10-20", "Re-analysis of saved D-043 predictions", "Earlier warning (A), pooled ranges (B), profile prior (C)",
         "Same 14 inspected people", "{wr.opened_at}", "No model refit", "Outer predictions read once",
         "{wr.b1.cov120}", "{wr.b0.cov120}", None, "B positive (not applied to serving); A and C negative",
         "A: +0.9 pt [−5.5, +6.6]; C: CI crosses zero", "reports/warning_range_experiment.md", "warnrange"),
        ("D-047", "—", "Software (public synthetic mode)", "Not a model experiment: opt-in public synthetic demonstration",
         "Synthetic participants only", None, "—", "not applicable", None, None, None, "Implemented; not deployed",
         "No performance claim", "docs/deploy.md", None),
        ("D-048", "gt-blend-2026-10-21", "Saved-prediction blend", "Does a LightGBM/CatBoost blend lower 60-min error by ≥ 0.5?",
         "12 development people (inner out-of-fold)", "{blend.dev_time}", "No fit (saved predictions)",
         "NOT OPENED (stop rule: target not met)", "{blend.inc60}-{blend.improvement_60}", "{blend.inc60}", "D-048",
         "Target not met; incumbent retained", "Gain 0.02 mg/dL vs needed 0.5", "docs/decisions.md D-048", "blend"),
        ("D-049", "gt-retrain-2026-10-21", "Retraining arms", "Do weighting or history features lower 60-min error by ≥ 0.5?",
         "12 development people (inner out-of-fold)", "{rt.started_at}", "24 fresh fits, fixed settings",
         "NOT OPENED (no arm met the gate)", None, None, "D-049", "B0 retained", "Best arm B1 gain 0.41 < 0.5; descriptive population-only result is after selection",
         "reports/retrain_study.md", "retrain"),
        ("D-050", "gt-ranges-2026-10-21", "Range shape study", "Does a trend-stratified range beat the global range before test?",
         "Serving pipeline; validation→calibration (pre-test), test read once", "{rg.generated_at}", "No model refit",
         "Final test read once, retrospective (inspected segment)", "{rg.pool.cov60}", "{rg.inc.cov60}", "D-050",
         "G kept; pooled calibration is a research candidate only", "Candidate coverage is on inspected data — not the served interval",
         "reports/serving_ranges_study.md", "ranges"),
        ("D-051", "gt-tune-2026-10-21", "Nested tuning", "Does tuned 24-h history (C1) beat equally tuned current features (C0)?",
         "14 people held out once (inspected)", "{tu.started_at}", "30 seeded trials per horizon (implemented budget)",
         "Outer folds opened once", "{tu.C1.mae60}", "{tu.C0.mae60}", "D-051", "Not adopted",
         "Small internal gain vs equally tuned C0 is not a release and is not the same as a comparison with the refitted serving recipe",
         "reports/tune_study.md", "tune"),
        ("D-052", "gt-joint-2026-10-21", "Joint model + correction selection", "Does joint selection on personalised error beat R0?",
         "14 people held out once (inspected)", "{jt.started_at}", "1,570 fits", "Outer folds opened once", "{jt.S.mae60}", "{jt.R0.mae60}",
         "D-052", "Not adopted", "High-target +1.70 > 1.0; warning guardrails failed; mixed-fit range erratum",
         "reports/joint_study.md", "joint"),
        ("D-053", "gt-constrained-2026-10-21", "Constrained (peak-aware) selection", "Can selection constrained on high-target error and workload beat R0?",
         "14 people held out once (inspected)", "{cs.started_at}", "1,490 logical / 2,930 physical fits", "Assessment opened once",
         "{cs.S.mae60}", "{cs.R0.mae60}", "D-053", "NOT ADOPTED (primary gate failed)",
         "High-target worsening +1.0658 vs +1.0 limit", "reports/constrained_study.md", "constrained"),
        ("D-054", "gt-foundation-2026-10-21", "Zero-shot foundation baseline", "Does a pretrained time-series model beat the tuned tree?",
         "14 people, all eligible origins (inspected)", "{fd.generated_at}", "No fit (zero-shot); one descriptive residual template",
         "Scored once", "{fd.FM60}", "{fd.R060}", None, "Not served; descriptive", "Not better than refitted recipe; worse above 180; pretraining data may include public CGM",
         "reports/foundation_baseline.md", "foundation"),
    ]
    for i, row_ in enumerate(studies, 5):
        dec, eid, kind, hyp, pop, start, budget, opening, cand, ctrl, gate, adopt, neg, report, fid = row_
        proto = f"configs/experiments/{eid}.yaml" if eid != "—" else "—"
        vals = [dec, eid, kind, hyp, pop, proto, commit_utc(proto) if proto != "—" else "—",
                xl_formula(start, ref) if start else "—", None, budget, opening,
                xl_formula(cand, ref) if cand else "—", xl_formula(ctrl, ref) if ctrl else "—", None,
                f"={gate_cells[gate]}" if gate else "—", adopt, neg, report,
                files[fid]["status"] if fid else "—"]
        for j, x in enumerate(vals, 1):
            st.cell(row=i, column=j, value=x)
        st.cell(row=i, column=9, value=f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),IF(G{i}<H{i},"commit precedes run","commit does NOT precede run"),"not applicable")')
        st.cell(row=i, column=14, value=f'=IF(AND(ISNUMBER(L{i}),ISNUMBER(M{i})),L{i}-M{i},"—")')
        for j in (7, 8):
            st.cell(row=i, column=j).number_format = "yyyy-mm-dd hh:mm:ss"
        for j in (12, 13, 14):
            st.cell(row=i, column=j).number_format = "0.00"
        if dec == "D-050":
            for j in (12, 13):
                st.cell(row=i, column=j).number_format = "0.0%"
            st.cell(row=i, column=14).number_format = "0.0%"
        if dec == "D-046":
            for j in (12, 13, 14):
                st.cell(row=i, column=j).number_format = "0.0%"
    finish(st, 4, len(cols))

    # ---- Verification
    ve = ws["Verification"]
    title(ve, "Verification — what was run, where, and with what outcome",
          "The latest full local matrix is dated to commit f073a10 (22 Oct); it is not a pass on any later revision. "
          "Rows marked author-reported have no saved log in the repository.")
    cols = [("Evidence class", 22), ("Run", 22), ("Suite", 34), ("Command (repo-relative)", 50), ("Tested commit", 14),
            ("Data mode", 18), ("Start (UTC)", 18), ("End (UTC)", 18), ("Exit", 6), ("Status", 10), ("Total", 8),
            ("Passed", 8), ("Failed", 8), ("Skipped", 8), ("Skip reasons", 40), ("Log", 40), ("JUnit", 40)]
    header(ve, 4, cols)
    i = 5

    def ledger_rows(fid: str, cls: str, run: str) -> None:
        nonlocal i
        led = load(FILES[fid][0])
        if not led:
            ve.cell(row=i, column=1, value=cls)
            ve.cell(row=i, column=2, value=run)
            ve.cell(row=i, column=3, value="unavailable")
            i += 1
            return
        for e in led.get("entries", []):
            cnt = e.get("counts") or {}
            mode = "recorded data" if "recorded" in (e.get("suite") or "") or "parity" in (e.get("suite") or "") else (
                "synthetic" if "public" in (e.get("suite") or "") else "local")
            vals = [cls, run, e.get("suite"), scrub(e.get("command") or ""), (e.get("revision") or {}).get("commit", "")[:10],
                    mode, utc((e.get("start") or {}).get("utc")), utc((e.get("end") or {}).get("utc")), e.get("exit_code"),
                    e.get("status"), cnt.get("total", "—"), cnt.get("passed", "—"), cnt.get("failed", "—"), cnt.get("skipped", "—"),
                    "; ".join(f"{k} ({v})" for k, v in (cnt.get("skip_reasons") or {}).items()), scrub(e.get("log") or ""),
                    scrub(e.get("junit") or "")]
            for j, x in enumerate(vals, 1):
                ve.cell(row=i, column=j, value=x)
            for j in (7, 8):
                ve.cell(row=i, column=j).number_format = "yyyy-mm-dd hh:mm:ss"
            i += 1

    ledger_rows("ledger", "Saved local result (final)", "final_2026-10-22_145326Z")
    for fid, run in (("parity", "reports/serving_parity.json"), ("parity_final", "final ledger parity copy")):
        vals = ["Saved local result (parity)", run, "App replay vs saved offline test forecasts", "uv run python scripts/serving_parity.py",
                "—", "recorded data", None, None, "—", f"={ref[f'{fid}.outcome']}",
                f"={ref[f'{fid}.offline_forecasts']}", f"={ref[f'{fid}.compared']}", f"={ref[f'{fid}.over_tolerance']}", "—",
                "tolerance 0.05 mg/dL (API displays 0.1 mg/dL)", FILES[fid][0], "—"]
        for j, x in enumerate(vals, 1):
            ve.cell(row=i, column=j, value=x)
        i += 1
    ledger_rows("ledger_overnight", "History (earlier ledger; failures preserved)", "overnight")
    for vals in (["Author-reported (logs not saved)", "fresh clone, 23 Oct", "pytest / vitest / e2e on a fresh clone", "make test; make test-web; make e2e",
                  "unknown", "synthetic", None, None, "—", "author-reported", "—", "vitest 141; e2e 16", "—", "e2e 13", "—", "not saved", "not saved"],
                 ["Author-reported (no saved CI log)", "GitHub Actions CI", "ruff, pytest, web jobs", ".github/workflows/ci.yml", "—", "synthetic",
                  None, None, "—", "author-reported", "—", "—", "—", "—", "—", "not saved", "—"]):
        for j, x in enumerate(vals, 1):
            ve.cell(row=i, column=j, value=x)
        i += 1
    finish(ve, 4, len(cols))

    # ---- Guardrails
    gu = ws["Guardrails"]
    title(gu, "Guardrails — non-numerical behaviour mapped to code and saved tests",
          "'tested' requires a named test that passed in the saved JUnit of the f073a10 ledger; code alone gives 'implemented'.")
    cols = [("Guardrail", 26), ("Claim", 50), ("Code / config", 44), ("Test(s)", 60), ("Saved result (JUnit, f073a10)", 22),
            ("Status", 14), ("Note", 36)]
    header(gu, 4, cols)
    junit = {}
    for fid in ("junit_pytest", "junit_e2e", "junit_public", "junit_summary"):
        p = ROOT / FILES[fid][0]
        if p.exists():
            for c in ET.parse(p).iter("testcase"):
                ok = not any(x.tag in ("failure", "error", "skipped") for x in c)
                junit[f"{c.get('classname')}::{c.get('name')}"] = ok
                junit[f"e2e: {c.get('name')}"] = ok

    def junit_status(tests: list[str]) -> tuple[str, str]:
        res = []
        for t in tests:
            key = t if t in junit else next((k for k in junit if k.endswith(t) or (t.startswith("e2e: ") and t[5:] in k)), None)
            res.append(junit.get(key) if key else None)
        if res and all(r is True for r in res):
            return "passed", "tested"
        if any(r is False for r in res):
            return "FAILED in saved run", "implemented"
        return "not found in saved JUnit", "implemented"

    guards = [
        ("Availability cutoff", "Only data available at the issue time is used (labs, meals, wearable, late readings)",
         "ml/glucotwin/features.py; ml/glucotwin/data/canonical.py",
         ["tests.test_leakage::test_future_values_cannot_change_earlier_features", "tests.test_leakage::test_late_logged_meal_is_hidden_until_available",
          "tests.test_data_expansion::test_a_reading_released_late_is_invisible_before_its_availability"]),
        ("Native readings only", "Interpolated 1-minute rows never count as observations", "ml/glucotwin/data/cgmacros.py",
         ["tests.test_leakage::test_one_minute_interpolated_rows_do_not_inflate_observations"]),
        ("Matured outcomes for correction", "The correction learns only from forecasts whose outcome is observed; never stacks",
         "ml/glucotwin/models/adaptation.py", ["tests.test_leakage::test_sixty_minute_residual_matures_only_when_target_is_available",
                                               "tests.test_leakage::test_corrections_never_stack_on_corrected_predictions"]),
        ("Bounded correction", "≥ 6 matured outcomes, ±25 mg/dL cap, reset after a 3-h gap", "ml/glucotwin/models/adaptation.py; configs/experiment.yaml",
         ["tests.test_replay_api::test_personalization_reports_its_bound_and_reset_gap"]),
        ("Abstention", "No forecast with insufficient history; nulls, not guesses", "ml/glucotwin/features.py",
         ["tests.test_replay_api::test_insufficient_history_abstains_with_nulls"]),
        ("Staleness + gap recovery", "A real sensor gap marks the forecast stale, the twin abstains, then recovers",
         "ml/glucotwin/features.py; services/api/glucotwin_api/replay.py",
         ["e2e: Participant 12: a real sensor gap, abstention without a made-up forecast, then recovery",
          "e2e: Participant 12: stale glucose before the gap closes, then abstention"]),
        ("Fusion needs its health record", "The profile model abstains until every profile value is available", "ml/glucotwin/features.py",
         ["tests.test_leakage::test_profile_model_is_gated_until_every_profile_value_is_available"]),
        ("Rewind / revision consistency", "Seeking back removes future state; forward replay equals reset + seek",
         "services/api/glucotwin_api/replay.py", ["tests.test_replay_api::test_seek_backward_removes_future_state_and_invalidates_cursors",
                                                   "tests.test_replay_api::test_forward_replay_equals_reset_and_seek"]),
        ("Immutable forecast identity", "A served forecast keeps its identity and value (parity with the saved offline forecast)",
         "services/api/glucotwin_api/replay.py; scripts/serving_parity.py",
         ["tests.test_v2_preserved::test_model_bundle_reproduces_saved_adaptive_forecasts[gt-primary-2026-10-13-fusion]"]),
        ("Frozen review snapshot", "A review summary is one session revision and clock; the copy stays frozen",
         "services/api/glucotwin_api/snapshot.py", ["tests.test_replay_api::test_review_snapshot_is_one_session_revision_and_clock",
                                                     "e2e: Participant 01, Day 9 20:38: a pending warning with a long review note; the copy stays frozen"]),
        ("Warning outcome pending", "A warning stays pending until its outcome is observable; gaps censor it",
         "ml/glucotwin/evaluation/alerts.py", ["tests.test_corrections::test_outcome_stays_pending_until_follow_up_is_observable_and_gaps_censor_it",
                                               "e2e: Participant 01: a warning is pending until its outcome is knowable"]),
        ("Visitor isolation (public mode)", "Each browser sees and controls only its own replays", "services/api/glucotwin_api/public.py",
         ["tests.test_public_demo::test_visitors_cannot_see_or_drive_each_others_sessions",
          "e2e: two simultaneous visitors cannot see or control each other's replays"]),
        ("Synthetic labelling", "Public mode serves synthetic participants only and labels them", "services/api/glucotwin_api/public.py",
         ["tests.test_public_demo::test_synthetic_only", "e2e: a fresh visitor sees the public synthetic demonstration, labelled as such"]),
        ("Test lock / no quiet re-run", "Editing a frozen protocol is refused; opened tests need a reason to re-run",
         "ml/glucotwin/research/assess.py", ["tests.test_research_pipeline::test_editing_a_frozen_protocol_is_refused"]),
        ("Unknown is not zero", "Zero onsets/warnings give unavailable metrics; unknown queue counts are never shown as zero",
         "ml/glucotwin/research/metrics.py", ["tests.test_research_units::test_zero_onsets_and_zero_warnings_give_unavailable_not_perfect_metrics",
                                              "e2e: queue: two replays are counted separately; an unknown count is never shown as zero"]),
        ("No treatment advice / no LLM predictions", "No insulin advice, triage or LLM-generated predictions", "apps/web; services/api (code inspection)",
         []),
    ]
    guard_status: dict[str, str] = {}
    for i, (name, claim, code, tests) in enumerate(guards, 5):
        res, stat = junit_status(tests) if tests else ("no test cited", "implemented")
        for j, x in enumerate([name, claim, code, "\n".join(tests) or "—", res, stat,
                               "Saved JUnit of the f073a10 ledger" if tests else "Absence of a feature; code inspection only"], 1):
            gu.cell(row=i, column=j, value=x)
        guard_status[name] = stat
    finish(gu, 4, len(cols))

    # ---- ClaimIndex
    claims = manifest_claims() + readme_claims() + frozen_slide_claims() + DECK_CLAIMS
    extra: dict[str, list[str]] = {}
    for cid, loc in EXTRA_LOCATIONS:
        extra.setdefault(cid, []).append(loc)
    ids = [c["id"] for c in claims]
    assert len(ids) == len(set(ids)), "duplicate claim id"
    gate_claim = {"C-50-05": "D-050", "C-51-08": "D-051", "C-52-05": "D-052", "C-53-12": "D-053 primary", "C-53-17": "D-053 operational"}
    gate_expect = {"C-50-05": "FAILED", "C-51-08": "FAILED", "C-52-05": "FAILED", "C-53-12": "FAILED", "C-53-17": "PASSED"}
    cl = ws_claims
    cl["A1"] = "ClaimIndex — every public claim, its evidence and status"
    cl["A1"].font = TITLE
    cl["A2"] = ("A source match is not clinical validation. Status counts below are live COUNTIF formulas. "
                "Workbook value = formula over Sources; Check = stated value within half a unit of its last decimal of the workbook value.")
    cl["A2"].font = NOTE
    nav = cl.cell(row=1, column=8, value="About / navigation →")
    nav.hyperlink = "#'About'!A1"
    nav.font = LINK
    HR = 6 + len(STATUSES)
    cols = [("Claim ID", 10), ("Exact claim", 54), ("Location(s)", 30), ("Evaluation / model scope", 34), ("Status", 18),
            ("Source IDs", 30), ("Workbook value (formula)", 13), ("Stated value", 11), ("Decimals", 8), ("Check", 12),
            ("Material limitation", 44), ("Unresolved action", 36), ("Manifest group | check", 36)]
    for k, s_ in enumerate(STATUSES):  # status name in B, live count in C, meaning in D
        r_ = 4 + k
        cl.cell(row=r_, column=2, value=s_).font = Font(bold=True, color=INK)
        cl.cell(row=r_, column=3, value=f'=COUNTIF($E${HR + 1}:$E${HR + len(claims)},B{r_})').font = Font(bold=True, color=BLUE)
        cl.cell(row=r_, column=4, value=STATUS_MEANING[s_]).font = NOTE
    cl.cell(row=3, column=2, value="Status").font = Font(bold=True, color=INK)
    cl.cell(row=3, column=3, value="Claims").font = Font(bold=True, color=INK)
    cl.cell(row=4 + len(STATUSES), column=2, value="total").font = Font(bold=True, color=INK)
    cl.cell(row=4 + len(STATUSES), column=3, value=f"=COUNTA($A${HR + 1}:$A${HR + len(claims)})").font = Font(bold=True, color=BLUE)
    header(cl, HR, cols)
    export = []
    for i, c in enumerate(claims, HR + 1):
        status = c["status"]
        pyv = None
        if c["expr"]:
            pyv = py_eval(c["expr"])
            if pyv is None:
                status = "pending"
            elif c["stated"] is not None and isinstance(c["stated"], (int, float)):
                if abs(pyv - float(c["stated"])) > 0.5 * 10 ** -c["dec"] + 1e-9:
                    status = "unsupported/contradicted"
                elif abs(abs(pyv - float(c["stated"])) - 0.5 * 10 ** -c["dec"]) < 1e-9:
                    c["limitation"] = (c["limitation"] + "; " if c["limitation"] else "") + (
                        f"Source value {pyv:g} is an exact rounding tie at the source's own precision; stated {c['stated']} "
                        "is one of its two valid roundings")
        if c["id"] in gate_claim:
            got = gate_py_out[gate_claim[c["id"]]]
            if got != gate_expect[c["id"]]:
                status = "unsupported/contradicted"
        if c["id"] == "C-EX-05" and pyv is not None and not (V.val("ext.cmp.lo60") <= 0 <= pyv):
            status = "unsupported/contradicted"
        tests = [t.strip() for t in c["sources"].split(";") if t.strip().startswith(("tests.", "e2e:"))]
        if status == "tested" and tests and junit_status(tests)[1] != "tested":
            status = "implemented"  # a cited test is missing from, or failed in, the saved JUnit
        if c["id"] == "C-S1-01":
            import yaml  # noqa: PLC0415

            team = (yaml.safe_load((ROOT / FILES["submission"][0]).read_text()) or {}).get("team", {})
            if not (team.get("name") == "Kryptonite" and team.get("leader") == "Avinash Gehi"
                    and "Hardik Hinduja" in (team.get("members") or []) and team.get("institution_short") == "VESIT"):
                status = "unsupported/contradicted"
        locs = c["locs"] + ("; " + "; ".join(extra[c["id"]]) if c["id"] in extra else "")
        srcs = "; ".join(dict.fromkeys(PH.findall(c["expr"] or ""))) + (("; " if c["expr"] else "") + c["sources"] if c["sources"] else "")
        cl.cell(row=i, column=1, value=c["id"])
        cl.cell(row=i, column=2, value=c["text"])
        cl.cell(row=i, column=3, value=locs)
        cl.cell(row=i, column=4, value=c["scope"])
        cl.cell(row=i, column=5, value=status)
        cl.cell(row=i, column=6, value=srcs)
        if c["id"] in gate_claim:
            cl.cell(row=i, column=7, value=f"={gate_cells[gate_claim[c['id']]]}")
            cl.cell(row=i, column=8, value=gate_expect[c["id"]])
            cl.cell(row=i, column=10, value=f'=IF(G{i}=H{i},"matches","MISMATCH")')
        elif c["expr"]:
            cl.cell(row=i, column=7, value=xl_formula(c["expr"], ref))
            if c["stated"] is not None:
                cl.cell(row=i, column=8, value=c["stated"])
                cl.cell(row=i, column=9, value=c["dec"])
                cl.cell(row=i, column=10, value=f'=IF(ISNUMBER(G{i}),IF(ABS(G{i}-H{i})<=0.5*10^(-I{i})+1E-9,"matches","MISMATCH"),"pending")')
            else:
                cl.cell(row=i, column=8, value="—")
                cl.cell(row=i, column=10, value="see limitation")
            cl.cell(row=i, column=7).number_format = "0.0000"
        else:
            for j in (7, 8, 10):
                cl.cell(row=i, column=j, value="—")
        cl.cell(row=i, column=11, value=c["limitation"] or "")
        cl.cell(row=i, column=12, value=c["action"] or ("" if status not in ("pending", "author-reported") else "See limitation"))
        cl.cell(row=i, column=13, value=c["manifest"])
        export.append({"id": c["id"], "claim": c["text"], "locations": locs, "scope": c["scope"], "status": status,
                       "source_ids": [x for x in srcs.split("; ") if x], "workbook_value": pyv, "stated": c["stated"],
                       "decimals": c["dec"], "limitation": c["limitation"], "action": c["action"], "manifest": c["manifest"]})
    finish(cl, HR, len(cols))
    cl.freeze_panes = cl.cell(row=HR + 1, column=3)

    # ---- About
    ab = ws["About"]
    ab["A1"] = "GlucoTwin research evidence workbook"
    ab["A1"].font = TITLE
    head = git("rev-parse", "--short", "HEAD")
    lines = [
        ("Built", dt.datetime.now(dt.UTC).strftime("%Y-%m-%d %H:%M UTC") + f" from repository HEAD {head} (working tree may differ)"),
        ("What it is", "A judge-readable evidence companion: every public claim → saved source value → visible calculation. "
                       "Not a new experiment; nothing was fitted or re-opened to build it."),
        ("Not", "Not clinical validation. A source match shows a number was reported faithfully, not that GlucoTwin helps patients."),
        ("Serving test wording", "Later recordings of the same 14 people (final test segment) — never 'unseen patients'."),
        ("Missing evidence", "'unavailable' / 'not opened' / 'pending' are shown as text, never as 0."),
        ("Provenance", "Most full artifacts are git-ignored (artifacts/): marked 'local / reproducible-only'. Values are stored "
                       "here with locators and SHA-256 so they can be checked after re-running the recorded-data path."),
        ("Rebuild", "uv run --group docs python scripts/build_research_workbook.py ; check: uv run --group docs python scripts/check_research_workbook.py"),
        ("Method note", "docs/submission/GlucoTwin_Research_Workbook_Methodology.md"),
    ]
    for k, (a, b_) in enumerate(lines, 3):
        ab.cell(row=k, column=1, value=a).font = Font(bold=True, color=INK)
        ab.cell(row=k, column=2, value=b_).alignment = WRAP
    r0 = 3 + len(lines) + 1
    ab.cell(row=r0, column=1, value="Sheets").font = Font(bold=True, color=BLUE)
    desc = {"ClaimIndex": "Primary view: claims, status, evidence, checks", "Sources": "Every source value once, with locator and hash",
            "Cohorts": "Who each result counts", "Forecasts": "Errors by evaluation, model and horizon (+ chart)",
            "WarningRanges": "Warning denominators, notice shares, range coverage (+ chart)",
            "Studies": "D-032, D-043 … D-054 ledger incl. negative and unopened results", "Gates": "Declared gates on unrounded values",
            "Verification": "Test runs, commits, logs; author-reported rows marked", "Guardrails": "Behaviour → code → saved test"}
    for k, (n, d_) in enumerate(desc.items(), r0 + 1):
        c = ab.cell(row=k, column=1, value=n)
        c.hyperlink = f"#'{n}'!A1"
        c.font = LINK
        ab.cell(row=k, column=2, value=d_)
    r1 = r0 + len(desc) + 2
    ab.cell(row=r1, column=1, value="Status meanings").font = Font(bold=True, color=BLUE)
    for k, s_ in enumerate(STATUSES, r1 + 1):
        ab.cell(row=k, column=1, value=s_)
        ab.cell(row=k, column=2, value=STATUS_MEANING[s_]).alignment = WRAP
    ab.column_dimensions["A"].width = 24
    ab.column_dimensions["B"].width = 120

    for w in wb.worksheets:
        w.sheet_properties.tabColor = BLUE if w.title == "ClaimIndex" else None
        w.sheet_view.showGridLines = w.title not in ("About",)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # no cached values are written; spreadsheet apps recalculate on open
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)

    counts = {s_: sum(1 for e in export if e["status"] == s_) for s_ in STATUSES}
    return {"claims": export, "counts": counts, "files": list(files.values()), "gates": gate_py_out,
            "sheet_rows": {w.title: w.max_row for w in wb.worksheets}, "guardrails": guard_status, "ref": ref}


# --------------------------------------------------------------------------- outputs

def write_json(res: dict) -> None:
    def clean(v: Any) -> Any:
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        if isinstance(v, (dt.datetime, dt.date)):
            return v.isoformat()
        return v

    out = {
        "generated_by": "scripts/build_research_workbook.py",
        "generated_at": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
        "repository_head": git("rev-parse", "--short", "HEAD"),
        "scope": "Aggregate claims and source values only; no participant rows. Serving test = later recordings of the same 14 people.",
        "status_counts": res["counts"],
        "gate_outcomes": res["gates"],
        "claims": [{k: clean(v) for k, v in c.items()} for c in res["claims"]],
        "source_files": res["files"],
        "source_values": [{"id": r["id"], "file": r["file"], "path": FILES[r["file"]][0], "locator": r["locator"],
                           "unit": r["unit"], "description": r["desc"],
                           "value": clean(r["value"]) if not isinstance(r["value"], (list, dict)) else json.dumps(r["value"])}
                          for r in V.rows],
    }
    OUT_JSON.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n")


METHOD = """# GlucoTwin research workbook — methodology and use

`GlucoTwin_Research_Workbook.xlsx` traces every public claim in the deck, README and narration to a saved source value
and shows the calculation. It is an evidence companion, not a new experiment: building it fits nothing and re-opens no
assessment. A source match shows a number was reported faithfully; it is **not** clinical validation.

## Sheets

| Sheet | Content |
|---|---|
| ClaimIndex | One row per claim: exact text, every location, scope, status, source IDs, a formula for the value, the stated value and a match check. Live status counts (COUNTIF) at the top. |
| Sources | Each source value stored **once**, at full precision, with its repository path, JSON pointer / field locator, source date, file SHA-256, protocol identity and publication status. All other sheets reference these cells. |
| Cohorts | Who each result counts: dataset, participants, split roles, horizon, origins, hours, onsets, inspected status. |
| Forecasts | MAE by evaluation group, model and horizon, persistence on the same origins, candidate − persistence and relative change (formulas), saved paired differences and CIs. Chart: serving vs persistence. |
| WarningRanges | Warning counts with explicit numerators and denominators, notice shares, false-alarm rates, time in warning, range coverage and per-person extremes; trend context with validation and final-test rows kept apart. Chart: range coverage vs the 80 % target. |
| Studies | D-032 and D-043 … D-054: purpose, population, protocol first-commit time vs run/open time, fit budget, opening status, candidate/control metric, gate outcome, adoption, negative findings. D-047 is software (public mode), not a model experiment. |
| Gates | The declared gates applied with formulas to unrounded values, next to each artifact's own saved pass/fail flag. |
| Verification | Test runs from the saved ledgers (commit, command, times, exit, counts, logs, JUnit), parity runs, and author-reported rows. |
| Guardrails | Non-numerical behaviours mapped to code and to tests that passed in the saved JUnit. |

## Status meanings

{statuses}

A numeric claim is checked automatically (and again by the `Check` formula): the stated value must lie within half a
unit of its last stated decimal of the source value, otherwise the builder sets the claim to `unsupported/contradicted`;
if the source file is absent it becomes `pending`. A source value that is itself an exact rounding tie at its own saved
precision (for example 22.305 saved with three decimals, stated as 22.30) is accepted and flagged in the limitation
column.

## Aggregation and conventions

* **MAE** is the mean absolute error in mg/dL, averaged **equally over participants** (participant-macro) unless a row
  says pooled. **RMSE** is reported where the source has it; it is not interchangeable with MAE.
* **Sign convention:** differences are **candidate − control**; negative means the candidate had lower error. Where a
  study saved a positive "gain" or "improvement" (D-048, D-049, D-050), the Gates sheet labels it as such.
* **Matching:** every difference compares the same forecast origins and people as its source analysis. Rows from
  different evaluation groups (serving test, participant-held-out internal, external CGM-only, synthetic) are never
  compared with each other.
* **Denominators:** notice shares (≥ 15 / 30 / 60 min) divide by **all** eligible onsets (64 on the serving test), so a
  missed onset counts against them; the median qualified lead uses **detected** onsets only (61). False alarms per
  observed patient-day divide the saved false-alarm count by the saved **observed** patient-day denominator; per
  eligible patient-day uses the different, smaller **eligible** base. Neither is recomputed from rounded hours.
* **Counts are different things:** 15,440 parity comparisons are forecasts (issued origins × two horizons); forecast
  origins, matched targets, participants and episodes are separate counts.
* **Uncertainty:** confidence intervals are the saved participant-cluster bootstrap endpoints from each artifact (seed
  and resample count as recorded by the study, e.g. 2,000 resamples for D-053). The workbook never creates a new
  interval and never treats overlapping windows as independent people.
* **Gates** use unrounded values: D-053's high-target worsening is about +1.0658 against a +1.0 limit, so the gate fails
  however the value is displayed.

## Populations

The serving model's results are **later recordings of the same 14 people** (final test segment, inspected,
retrospective). Participant-held-out studies reuse 14 already-inspected people (internal validation, not an untouched
cohort). The 98-person ShanghaiT2DM checks test a **CGM-only** procedure, not the serving fusion model. CGMacros has 45
participants; 14 form the Type 2 evaluation group. Synthetic participants are fictional and never evidence.

## Freshness and provenance

Each Sources row carries the source file's SHA-256 and date (the artifact's own timestamp where it has one, else the
last git commit, else the file time). Most full artifacts live under `artifacts/`, which is **git-ignored**: those rows
are marked `local / reproducible-only (git-ignored)` and a judge cannot download them from GitHub; they can be
regenerated with the recorded-data path (`make data-download && make data && make experiment`, then the research
targets), after which the hashes can be compared. Files marked `public in repo (tracked)` are in the repository.
Protocol-before-run uses git commit times of `configs/experiments/*.yaml`; that is local commit evidence, not a
clinical trial registration.

The latest full local test matrix is dated to commit `f073a10` (22 Oct 2026); it is not a pass on any later revision.
Fresh-clone checks reported by the team on 23 Oct have no saved logs and are marked author-reported.

## Rebuild and check

```bash
uv run --group docs python scripts/build_research_workbook.py
uv run --group docs python scripts/check_research_workbook.py   # LibreOffice recalculation + checks
```

To add slide or narration claims, edit `DECK_CLAIMS` / `EXTRA_LOCATIONS` in `scripts/build_research_workbook.py` and
rebuild.
"""


def main() -> int:
    res = build()
    write_json(res)
    OUT_MD.write_text(METHOD.replace("{statuses}", "\n".join(f"* **{s}** — {STATUS_MEANING[s]}" for s in STATUSES)))
    print(f"wrote {OUT_XLSX.relative_to(ROOT)}, {OUT_JSON.relative_to(ROOT)}, {OUT_MD.relative_to(ROOT)}")
    print("status counts:", json.dumps(res["counts"]))
    print("gates:", json.dumps(res["gates"]))
    print("sheet rows:", json.dumps(res["sheet_rows"]))
    bad = [c["id"] for c in res["claims"] if c["status"] == "unsupported/contradicted"]
    if bad:
        print("CONTRADICTED:", ", ".join(bad))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
